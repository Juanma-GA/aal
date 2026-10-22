"""The generic Excel engine (E03, app/io_excel) on synthetic workbooks:
headers of one and two rows, every column type and its normalization,
errors (length, obligation, value out of its table, formula without a
value), the file checks (.xlsm, sizes, macros, corrupt) and the export's
cell formats."""

import io
import zipfile
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.io_excel.analysis import Stored, analyze
from app.io_excel.cells import CellError, normalize
from app.io_excel.reader import ExcelFileError, Limits, check_upload, read_rows
from app.io_excel.template import Column, ColumnType, Template
from app.io_excel.writer import write_workbook
from excel_compare import workbook_differences

LIMITS = Limits(max_bytes=1024 * 1024, max_uncompressed_bytes=10 * 1024 * 1024, max_rows=100)

ONE = Template(
    code="test",
    sheet="Datos",
    columns=(
        Column("code", ("Código",), max_length=2, required=True, key=True, zero_pad=2),
        Column("name", ("Nombre",), max_length=10, required=True),
        Column("count", ("Cantidad",), ColumnType.INTEGER, max_length=3),
        Column("price", ("Precio",), ColumnType.DECIMAL, max_length=6, decimals=2),
        Column("flag", ("Reparable",), ColumnType.YES_NO),
        Column("since", ("Desde",), ColumnType.DATE),
        Column("unit", ("Unidad",), ref_table="D112"),
        Column("kind", ("Tipo",), choices=(("temporal", "Temporal"), ("no_temporal", "No temporal"))),
    ),
)

TWO = Template(
    code="test2",
    sheet="Anexo",
    header_rows=2,
    columns=(
        Column("code", ("Identificación", "Código"), key=True, required=True),
        Column("name", ("Identificación", "Nombre")),
        Column("qty", ("Datos", "Cantidad"), ColumnType.INTEGER),
    ),
)


def book(rows: list[list], sheet: str = "Datos") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


HEADER = ["Código", "Nombre", "Cantidad", "Precio", "Reparable", "Desde", "Unidad", "Tipo"]


def col(name: str) -> Column:
    return next(c for c in ONE.columns if c.field == name)


# ── normalization, one rule at a time ──────────────────────────────────────


@pytest.mark.parametrize(
    "raw, expected",
    [(2, "02"), (2.0, "02"), ("  AB ", "AB"), ("7", "7"), (12, "12")],
)
def test_text_codes(raw, expected):
    assert normalize(col("code"), raw).value == expected


def test_a_padded_number_is_a_warning():
    assert normalize(col("code"), 2).warning == ("number_read_as_code", {"value": "2", "code": "02"})
    assert normalize(col("code"), "02").warning is None


@pytest.mark.parametrize("raw, code", [("ABC", "too_long"), (None, "required"), ("   ", "required"), (True, "not_text")])
def test_text_errors(raw, code):
    with pytest.raises(CellError) as exc:
        normalize(col("code"), raw)
    assert exc.value.code == code


def test_too_long_is_never_cut():
    with pytest.raises(CellError) as exc:
        normalize(col("name"), "x" * 11)
    assert exc.value.params == {"length": 11, "max": 10}


@pytest.mark.parametrize("raw, expected", [(5, 5), (5.0, 5), (" 12 ", 12), ("-3", -3), (None, None)])
def test_integers(raw, expected):
    assert normalize(col("count"), raw).value == expected


@pytest.mark.parametrize("raw, code", [(2.5, "not_integer"), ("doce", "not_integer"), (True, "not_integer"), (1234, "too_many_digits")])
def test_integer_errors(raw, code):
    with pytest.raises(CellError) as exc:
        normalize(col("count"), raw)
    assert exc.value.code == code


@pytest.mark.parametrize("raw, expected", [(3.5, Decimal("3.50")), ("3,25", Decimal("3.25")), ("10", Decimal("10.00")), (7, Decimal("7.00"))])
def test_decimals(raw, expected):
    value = normalize(col("price"), raw).value
    assert value == expected and str(value) == str(expected)


@pytest.mark.parametrize("raw, code", [(1.234, "too_many_decimals"), ("1.234,5", "not_decimal"), ("abc", "not_decimal"), (12345.6, "too_many_digits")])
def test_decimal_errors_never_round(raw, code):
    with pytest.raises(CellError) as exc:
        normalize(col("price"), raw)
    assert exc.value.code == code


@pytest.mark.parametrize("raw, expected", [("s", True), ("N", False), (" S ", True), (True, True), (False, False)])
def test_yes_no(raw, expected):
    assert normalize(col("flag"), raw).value is expected


def test_yes_no_error():
    with pytest.raises(CellError) as exc:
        normalize(col("flag"), "Sí")
    assert exc.value.code == "not_yes_no"


@pytest.mark.parametrize("raw", [20261022, "20261022", 20261022.0, date(2026, 10, 22)])
def test_dates(raw):
    assert normalize(col("since"), raw).value == date(2026, 10, 22)


@pytest.mark.parametrize("raw", ["2026-10-22", 20261340, "2026102"])
def test_date_errors(raw):
    with pytest.raises(CellError) as exc:
        normalize(col("since"), raw)
    assert exc.value.code == "not_a_date"


def test_choices_by_code_or_label_any_case():
    assert normalize(col("kind"), "No temporal").value == "no_temporal"
    assert normalize(col("kind"), "TEMPORAL").value == "temporal"
    assert normalize(col("kind"), "no_temporal").value == "no_temporal"
    with pytest.raises(CellError) as exc:
        normalize(col("kind"), "otro")
    assert exc.value.params == {"allowed": "Temporal, No temporal"}


def test_excel_escapes_are_undone():
    assert normalize(col("name"), "a_x000D_b").value == "a\rb"


# ── reading ────────────────────────────────────────────────────────────────


def test_one_header_row_and_the_values():
    data = book([HEADER, [2, " Tornillo ", 5, 3.5, "s", 20261022, "EA", "Temporal"], [], ["AB", "Tuerca"]])
    read = read_rows(ONE, data, LIMITS)
    assert read.errors == []
    assert [r.row for r in read.rows] == [2, 4]  # the empty row 3 is skipped
    first = read.rows[0]
    assert first.values == {
        "code": "02", "name": "Tornillo", "count": 5, "price": Decimal("3.50"), "flag": True,
        "since": date(2026, 10, 22), "unit": "EA", "kind": "temporal",
    }
    assert [w.code for w in first.warnings] == ["number_read_as_code"]


def test_a_wrong_header_stops_the_reading():
    read = read_rows(ONE, book([["Código", "Nombre del artículo", *HEADER[2:]], ["AB", "x"]]), LIMITS)
    assert [(i.code, i.column, i.params) for i in read.errors] == [
        ("header_mismatch", "B", {"expected": "Nombre", "found": "Nombre del artículo"})
    ]
    assert read.rows == []


def test_a_missing_sheet():
    read = read_rows(ONE, book([HEADER], sheet="Otra"), LIMITS)
    assert [i.code for i in read.errors] == ["sheet_missing"]


def test_extra_columns_are_a_warning():
    read = read_rows(ONE, book([HEADER + ["Notas"], ["AB", "x"]]), LIMITS)
    assert read.errors == [] and [i.code for i in read.warnings] == ["extra_column_ignored"]


def test_two_header_rows_with_merged_groups():
    exported = write_workbook(TWO, [{"code": "A1", "name": "Uno", "qty": 3}])
    sheet = openpyxl.load_workbook(io.BytesIO(exported))["Anexo"]
    assert sorted(map(str, sheet.merged_cells.ranges)) == ["A1:B1"]
    read = read_rows(TWO, exported, LIMITS)
    assert read.errors == [] and read.rows[0].values == {"code": "A1", "name": "Uno", "qty": 3}
    # A wrong group is caught in the first row, a wrong title in the second.
    wrong = book([["Identificación", None, "Otros"], ["Código", "Nombre", "Cantidad"]], sheet="Anexo")
    assert [(i.row, i.column) for i in read_rows(TWO, wrong, LIMITS).errors] == [(1, "C")]


def test_errors_by_row_and_column():
    data = book([HEADER, ["ABC", None, "x"], ["AB", "y" * 11]])
    read = read_rows(ONE, data, LIMITS)
    found = [(e.row, e.column, e.code) for r in read.rows for e in r.errors]
    assert found == [(2, "A", "too_long"), (2, "B", "required"), (2, "C", "not_integer"), (3, "B", "too_long")]


def test_a_formula_without_a_stored_value_is_an_error():
    # openpyxl writes formulas without their cached value, like a file
    # never recalculated.
    data = book([HEADER, ["AB", "=CONCAT(\"a\",\"b\")"]])
    read = read_rows(ONE, data, LIMITS)
    assert [(e.column, e.code) for e in read.rows[0].errors] == [("B", "formula_without_value")]


def test_duplicate_keys():
    read = read_rows(ONE, book([HEADER, ["AB", "x"], ["AB", "y"]]), LIMITS)
    assert [(e.row, e.code, e.params) for r in read.rows for e in r.errors] == [(3, "duplicate_key", {"first_row": 2})]


def test_too_many_rows():
    with pytest.raises(ExcelFileError) as exc:
        read_rows(ONE, book([HEADER] + [[f"{i:02d}", "x"] for i in range(5)]), Limits(LIMITS.max_bytes, LIMITS.max_uncompressed_bytes, 4))
    assert exc.value.code == "excel_too_many_rows"


# ── file checks ────────────────────────────────────────────────────────────


@pytest.mark.parametrize("name", ["datos.xlsm", "datos.xls", "datos.csv", None])
def test_only_xlsx(name):
    with pytest.raises(ExcelFileError) as exc:
        check_upload(name, book([HEADER]), LIMITS)
    assert exc.value.code == "excel_wrong_extension"


def test_an_xlsx_with_macros_inside_is_refused():
    buf = io.BytesIO(book([HEADER]))
    with zipfile.ZipFile(buf, "a") as z:
        z.writestr("xl/vbaProject.bin", b"\x00" * 10)
    with pytest.raises(ExcelFileError) as exc:
        check_upload("datos.xlsx", buf.getvalue(), LIMITS)
    assert exc.value.code == "excel_has_macros"


def test_size_limits():
    data = book([HEADER])
    with pytest.raises(ExcelFileError) as exc:
        check_upload("datos.xlsx", data, Limits(len(data) - 1, LIMITS.max_uncompressed_bytes, 10))
    assert exc.value.code == "excel_too_large"
    # A zip bomb: small file, large content -- refused from the zip's own
    # directory, never expanded.
    buf = io.BytesIO(data)
    with zipfile.ZipFile(buf, "a", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/media/padding.bin", b"\x00" * (3 * 1024 * 1024))
    bomb = buf.getvalue()
    assert len(bomb) < 1024 * 1024
    with pytest.raises(ExcelFileError) as exc:
        check_upload("datos.xlsx", bomb, Limits(LIMITS.max_bytes, 2 * 1024 * 1024, 10))
    assert exc.value.code == "excel_too_large_uncompressed"


@pytest.mark.parametrize("data, code", [(b"", "excel_empty"), (b"not a zip", "excel_not_xlsx")])
def test_empty_or_not_a_workbook(data, code):
    with pytest.raises(ExcelFileError) as exc:
        check_upload("datos.xlsx", data, LIMITS)
    assert exc.value.code == code


def test_a_corrupt_workbook():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/workbook.xml", "<roto")
    check_upload("datos.xlsx", buf.getvalue(), LIMITS)
    with pytest.raises(ExcelFileError) as exc:
        read_rows(ONE, buf.getvalue(), LIMITS)
    assert exc.value.code == "excel_corrupt"


# ── export ─────────────────────────────────────────────────────────────────


def test_export_keeps_codes_as_text_and_reads_back_the_same():
    rows = [
        {"code": "02", "name": "=SUMA(1)", "count": 5, "price": Decimal("3.50"), "flag": True, "since": date(2026, 10, 22), "kind": "no_temporal"},
        {"code": "AB", "name": "_x000D_ x"},
    ]
    data = write_workbook(ONE, rows)
    sheet = openpyxl.load_workbook(io.BytesIO(data))["Datos"]
    assert [c.value for c in sheet[1]] == HEADER
    a2 = sheet["A2"]
    assert (a2.value, a2.data_type, a2.number_format) == ("02", "s", "@")
    assert (sheet["B2"].value, sheet["B2"].data_type) == ("=SUMA(1)", "s")  # never a formula
    assert (sheet["C2"].value, sheet["C2"].number_format) == (5, "0")
    assert (sheet["D2"].value, sheet["D2"].number_format) == (3.5, "0.00")
    assert [sheet["E2"].value, sheet["F2"].value, sheet["H2"].value] == ["S", "20261022", "No temporal"]
    assert sheet.column_dimensions["A"].number_format == "@"
    read = read_rows(ONE, data, LIMITS)
    assert read.errors == [] and all(not r.errors for r in read.rows)
    assert read.rows[0].values["name"] == "=SUMA(1)" and read.rows[1].values["name"] == "_x000D_ x"
    assert read.rows[0].values["price"] == Decimal("3.50") and read.rows[0].values["since"] == date(2026, 10, 22)
    assert workbook_differences(data, write_workbook(ONE, rows)) == []


def test_the_compare_helper_finds_a_changed_cell():
    rows = [{"code": "AB", "name": "Uno"}]
    other = [{"code": "AB", "name": "Dos"}]
    assert workbook_differences(write_workbook(ONE, rows), write_workbook(ONE, other)) == ["Datos!B2 value: 'Uno' != 'Dos'"]


# ── analysis ───────────────────────────────────────────────────────────────


def test_analysis_new_modified_unchanged_and_references():
    data = book([HEADER, ["AA", "Igual", None, None, None, None, "EA"], ["BB", "Cambiado"], ["CC", "Nuevo", None, None, None, None, "ZZ"], ["DD", "Viejo", None, None, None, None, "OLD"], ["EE", "Ya", None, None, None, None, "OLD"]])
    stored = {
        ("AA",): Stored({"code": "AA", "name": "Igual", "unit": "EA"}, "1", 1),
        ("BB",): Stored({"code": "BB", "name": "Antes"}, "2", 4),
        ("EE",): Stored({"code": "EE", "name": "Ya", "unit": "OLD"}, "5", 1),
    }
    references = {"D112": {"EA": True, "OLD": False}}
    report = analyze(ONE, read_rows(ONE, data, LIMITS), stored, references)
    assert report["counts"] == {"rows": 5, "new": 0, "modified": 1, "unchanged": 2, "errors": 2, "rows_with_errors": 2, "warnings": 0}
    assert report["modified"] == [{"row": 3, "key": "BB", "changes": [{"field": "name", "header": "Nombre", "before": "Antes", "after": "Cambiado"}]}]
    assert report["pending"]["modified"] == [{"row": 3, "ref": "2", "version": 4, "values": {"code": "BB", "name": "Cambiado", "count": None, "price": None, "flag": None, "since": None, "unit": None, "kind": None}}]
    # ZZ is not in the table; OLD is deactivated: refused for a new value,
    # kept for the value that already had it (row 6).
    assert [(e["row"], e["code"]) for e in report["errors"]] == [(4, "not_in_table"), (5, "inactive_in_table")]
    assert {u["key"] for u in report["unchanged"]} == {"AA", "EE"}
