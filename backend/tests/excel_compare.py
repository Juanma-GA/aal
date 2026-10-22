"""Test helper (E03): two workbooks compared cell by cell -- sheets and
their order, merged ranges, and for every cell its value, type and number
format. Returns the differences, empty when the books are the same."""

import io

import openpyxl


def workbook_differences(a: bytes, b: bytes) -> list[str]:
    book_a = openpyxl.load_workbook(io.BytesIO(a))
    book_b = openpyxl.load_workbook(io.BytesIO(b))
    if book_a.sheetnames != book_b.sheetnames:
        return [f"sheets {book_a.sheetnames} != {book_b.sheetnames}"]
    found = []
    for name in book_a.sheetnames:
        sa, sb = book_a[name], book_b[name]
        if sorted(map(str, sa.merged_cells.ranges)) != sorted(map(str, sb.merged_cells.ranges)):
            found.append(f"{name}: merged cells differ")
        rows = max(sa.max_row, sb.max_row)
        cols = max(sa.max_column, sb.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = sa.cell(r, c), sb.cell(r, c)
                for what in ("value", "data_type", "number_format"):
                    if getattr(ca, what) != getattr(cb, what):
                        found.append(f"{name}!{ca.coordinate} {what}: {getattr(ca, what)!r} != {getattr(cb, what)!r}")
    return found
