"""The NT 10/17 validator (E16) over the demo: the figures of the encargo
§2.3, the rules' configuration (admin, audited, "desactualizado"),
accepting a warning, the delivery's block, the reports (Excel, PDF), the
ESTADO DE VALIDACIÓN sheet, the permissions, and the time with 714 rows
(validating, and analysing and applying a book)."""

import asyncio
import io
import time
import uuid
import zipfile
from collections import Counter

import openpyxl
from sqlalchemy import select, update

from app.db.base import async_session_factory
from app.models import AuditLog, Material, PhysicalPosition, ValidationRule
from integration.test_anexo_xi import _jobs, analysed, apply  # noqa: F401
from integration.test_catalog import catalog, imported  # noqa: F401
from integration.test_confunfis import PUESTO, TRAMONTANA, all_rows, base, demo, find  # noqa: F401

# What the demo has (demo/LEEME_DATOS_DEMO.md, encargo §2.3).
DEMO_FIGURES = {
    "CF-001": 1, "CF-002": 1, "CF-007": 1, "CF-012": 1, "CF-014": 1, "CF-015": 1, "CF-017": 1, "PF-001": 1,
    "AR-001": 1, "AR-002": 1, "AR-003": 1, "AR-006": 1, "AR-007": 1,
}


def url(demo, name=TRAMONTANA):
    return f"/api/materials/{demo['materials'][name].id}/validation"


async def validate(client, demo, headers, name=TRAMONTANA):
    res = await client.post(url(demo, name), headers=headers)
    assert res.status_code == 200, res.text
    return res.json()


async def _restore_rules(codes):
    from app.validacion.catalog import BY_CODE

    async with async_session_factory() as db:
        for code in codes:
            await db.execute(update(ValidationRule).where(ValidationRule.code == code).values(active=True, severity=BY_CODE[code].severity, params=BY_CODE[code].params, help=None))
        await db.commit()


async def test_the_figures_of_the_demo(client, demo):
    analyst = demo["people"]["analista"]
    assert (await client.get(url(demo), headers=analyst)).json()["state"] == "none"
    report = await validate(client, demo, analyst)
    counts = Counter(f["rule_code"] for f in report["findings"])
    assert dict(counts) == DEMO_FIGURES
    by_rule = {r["code"]: r for r in report["by_rule"]}
    blocking_cf = sorted(c for c, r in by_rule.items() if c.startswith("CF-") and r["severity"] == "bloqueante")
    assert blocking_cf == ["CF-001", "CF-002", "CF-007", "CF-012", "CF-014", "CF-015"]
    assert {c for c, r in by_rule.items() if r["severity"] == "aviso"} == {"CF-017", "AR-001", "AR-007"}
    assert (report["state"], report["blocking"], report["warnings"], report["outdated"]) == ("red", 10, 3, False)
    assert report["seconds"] < 3
    # Each one says where it is and opens a row.
    air = next(f for f in report["findings"] if f["rule_code"] == "CF-007")
    assert "A24110 01 00" in air["location"] and "FILTRO" in air["location"].upper() and air["field"] == "cs" and air["row_id"]
    tyre = next(f for f in report["findings"] if f["rule_code"] == "CF-017")
    assert "261033X900805" in str(tyre["params"]) and tyre["row_id"]
    # The CONFUNFIS's chip: n errores · n avisos (by the rules' severity).
    summary = (await client.get(f"{base(demo)}/summary", headers=analyst)).json()
    assert (summary["errors"], summary["warnings"]) == (6, 1)


async def test_a_change_makes_it_outdated_and_the_fix_clears_it(client, demo):
    analyst = demo["people"]["analista"]
    await validate(client, demo, analyst)
    [air] = find((await all_rows(client, demo, analyst))["items"], "A24110")
    res = await client.patch(f"{base(demo)}/rows/{air['id']}", json={"version": air["version"], "cs_sns": "A20000", "cs_co": "01", "cs_al": "00"}, headers=analyst)
    assert res.status_code == 200, res.text
    stale = (await client.get(url(demo), headers=analyst)).json()
    assert stale["outdated"] and stale["outdated_data"] and not stale["outdated_rules"]
    again = await validate(client, demo, analyst)
    assert "CF-007" not in {f["rule_code"] for f in again["findings"]} and again["blocking"] == 9 and not again["outdated"]


async def test_the_rules_configuration(client, demo):
    admin, analyst = demo["people"]["admin"], demo["people"]["analista"]
    rules = {r["code"]: r for r in (await client.get("/api/validation-rules", headers=analyst)).json()}
    assert len(rules) == 30 and rules["CF-018"]["severity"] == "aviso" and rules["CF-001"]["ref"]
    body = {"version": rules["CF-017"]["version"], "severity": "bloqueante"}
    assert (await client.patch("/api/validation-rules/CF-017", json=body, headers=analyst)).status_code == 403
    await validate(client, demo, analyst)
    async with async_session_factory() as db:
        logged_before = len((await db.execute(select(AuditLog.id).where(AuditLog.action == "validation.rule_updated", AuditLog.entity_id == "CF-017"))).all())
    try:
        res = await client.patch("/api/validation-rules/CF-017", json=body, headers=admin)
        assert res.status_code == 200 and res.json()["severity"] == "bloqueante"
        assert (await client.patch("/api/validation-rules/CF-017", json=body, headers=admin)).status_code == 409
        bad = await client.patch("/api/validation-rules/CF-019", json={"version": rules["CF-019"]["version"], "params": {"nope": 1}}, headers=admin)
        assert (bad.status_code, bad.json()["detail"]["code"]) == (422, "validation_rule_param_unknown")
        assert (await client.patch("/api/validation-rules/XX-999", json=body, headers=admin)).status_code == 404
        stale = (await client.get(url(demo), headers=analyst)).json()
        assert stale["outdated"] and stale["outdated_rules"]
        report = await validate(client, demo, analyst)
        assert report["blocking"] == 11 and report["warnings"] == 2
        # Off: not reported at all.
        off = await client.patch("/api/validation-rules/CF-017", json={"version": res.json()["version"], "active": False}, headers=admin)
        assert off.status_code == 200
        assert "CF-017" not in {f["rule_code"] for f in (await validate(client, demo, analyst))["findings"]}
        async with async_session_factory() as db:
            logged = (await db.execute(
                select(AuditLog).where(AuditLog.action == "validation.rule_updated", AuditLog.entity_id == "CF-017").order_by(AuditLog.created_at)
            )).scalars().all()[logged_before:]
        assert len(logged) == 2 and logged[0].before["severity"] == "aviso" and logged[0].after["severity"] == "bloqueante"
    finally:
        await _restore_rules(["CF-017"])


async def test_accepting_a_warning(client, demo):
    analyst, observer = demo["people"]["analista"], demo["people"]["observador"]
    report = await validate(client, demo, observer)  # validating is confunfis.ver
    tyre = next(f for f in report["findings"] if f["rule_code"] == "CF-017")
    blocking = next(f for f in report["findings"] if f["severity"] == "bloqueante")
    accept = f"{url(demo)}/findings/{tyre['id']}/accept"
    body = {"justification": "Neumático también en cadenas: error de demo conocido"}
    assert (await client.post(accept, json=body, headers=observer)).status_code == 403
    assert (await client.post(accept, json={"justification": "  "}, headers=analyst)).status_code == 422
    res = await client.post(accept, json=body, headers=analyst)
    assert res.status_code == 200, res.text
    after = res.json()
    mine = next(f for f in after["findings"] if f["id"] == tyre["id"])
    assert mine["accepted"] and mine["justification"] == body["justification"] and mine["accepted_by"]
    assert (after["warnings"], after["accepted"]) == (2, 1)
    again = await client.post(accept, json=body, headers=analyst)
    assert (again.status_code, again.json()["detail"]["code"]) == (409, "finding_already_accepted")
    refused = await client.post(f"{url(demo)}/findings/{blocking['id']}/accept", json=body, headers=analyst)
    assert (refused.status_code, refused.json()["detail"]["code"]) == (409, "finding_blocking_cannot_be_accepted")
    # It holds in the next runs.
    next_run = await validate(client, demo, analyst)
    assert next((f for f in next_run["findings"] if f["rule_code"] == "CF-017"))["accepted"] and next_run["accepted"] == 1
    assert (await client.post(f"{url(demo)}/findings/{uuid.uuid4()}/accept", json=body, headers=analyst)).status_code == 404


async def test_the_delivery_is_blocked_until_it_validates_green(client, demo):
    analyst = demo["people"]["analista"]
    delivery = f"{base(demo)}/delivery"
    res = await client.get(delivery, headers=analyst)
    assert (res.status_code, [p["code"] for p in res.json()["detail"]["problems"]]) == (409, ["delivery_not_validated"])
    await validate(client, demo, analyst)
    res = await client.get(delivery, headers=analyst)
    assert [p["code"] for p in res.json()["detail"]["problems"]] == ["delivery_blocking"]
    # Every rule a warning: no blocking findings, green enough to deliver.
    async with async_session_factory() as db:
        codes = list((await db.execute(select(ValidationRule.code))).scalars())
        await db.execute(update(ValidationRule).values(severity="aviso"))
        await db.commit()
    try:
        await validate(client, demo, analyst)
        res = await client.get(delivery, headers=analyst)
        assert res.status_code == 200 and res.headers["content-type"] == "application/zip"
        names = zipfile.ZipFile(io.BytesIO(res.content)).namelist()
        assert sorted(n.split("_")[0] for n in names) == ["ANEXO", "FICHERO"]
        anexo = openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(res.content)).read(next(n for n in names if n.startswith("ANEXO")))))
        assert "ESTADO DE VALIDACIÓN" not in anexo.sheetnames and anexo.sheetnames[0] == "CON_FUN_FIS"
        # A change afterwards: outdated, refused again.
        [air] = find((await all_rows(client, demo, analyst))["items"], "A24110")
        await client.patch(f"{base(demo)}/rows/{air['id']}", json={"version": air["version"], "qty": 2}, headers=analyst)
        res = await client.get(delivery, headers=analyst)
        assert [p["code"] for p in res.json()["detail"]["problems"]] == ["delivery_outdated"]
    finally:
        await _restore_rules(codes)


async def test_the_status_sheet_and_the_reports(client, demo):
    analyst = demo["people"]["analista"]
    assert (await client.get(f"{url(demo)}/report", headers=analyst)).status_code == 409
    await validate(client, demo, analyst)
    book = openpyxl.load_workbook(io.BytesIO((await client.get(f"{base(demo)}/export", params={"kind": "aal"}, headers=analyst)).content))
    status = book["ESTADO DE VALIDACIÓN"]
    assert book.sheetnames[0] == "ESTADO DE VALIDACIÓN"
    assert status["A1"].value == TRAMONTANA and status["B1"].value.startswith("BORRADOR") and "10" in status["B1"].value and "3" in status["B1"].value
    pdf = await client.get(f"{url(demo)}/report", params={"format": "pdf"}, headers=analyst)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and "VALIDACION_NT" in pdf.headers["content-disposition"]
    xlsx = await client.get(f"{url(demo)}/report", params={"format": "xlsx"}, headers=analyst)
    report = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert report.sheetnames == ["RESUMEN", "HALLAZGOS"]
    rows = list(report["HALLAZGOS"].iter_rows(values_only=True))
    assert len(rows) == 1 + sum(DEMO_FIGURES.values()) and any("CF-007" in (r[0] or "") for r in rows[1:])
    assert (await client.get(f"{url(demo)}/report", params={"format": "doc"}, headers=analyst)).status_code == 422


async def test_permissions(client, demo):
    outsider = demo["people"]["outsider"]
    # A CUF the user cannot open does not exist for them.
    assert (await client.get(url(demo), headers=outsider)).status_code == 404
    assert (await client.post(url(demo), headers=outsider)).status_code == 404
    assert (await client.get(f"{base(demo)}/delivery", headers=outsider)).status_code == 404


async def test_applying_a_book_validates_the_cuf(client, demo):
    analyst = demo["people"]["analista"]
    data = (await client.get(f"{base(demo)}/export", params={"kind": "anexo"}, headers=analyst)).content
    job = await analysed(client, demo, analyst, data)
    assert job["report"]["result"]["blocking"] == 10 and job["report"]["result"]["severity"]["CF-017"] == "aviso"
    res = await apply(client, demo, analyst, job)
    assert res.status_code == 200, res.text
    latest = (await client.get(url(demo), headers=analyst)).json()
    assert latest["trigger"] == "import" and latest["blocking"] == 10


async def test_714_rows_are_fast(client, demo):
    """The size of the real book: validated in < 3 s; its book analysed and applied in bounded time."""
    analyst = demo["people"]["analista"]
    material = demo["materials"][TRAMONTANA]
    async with async_session_factory() as db:
        mid = (await db.get(Material, material.id)).id
        rows = (await db.execute(select(PhysicalPosition).where(PhysicalPosition.material_id == mid, PhysicalPosition.deleted_at.is_(None)))).scalars().all()
        template = next(r for r in rows if r.sns == "A21031")
        for n in range(714 - len(rows)):
            db.add(PhysicalPosition(
                id=uuid.uuid4(), material_id=mid, sns="A21031", co=f"{2 + n // 99:02d}", al="00", qty=1,
                article_name=f"PIEZA {n}", function_text="Prueba", ct_sigle="N", cs_sns="A21030", cs_co="01", cs_al="00", source="manual", version=1,
                article_id=template.article_id,
            ))
        await db.commit()
    assert (await all_rows(client, demo, analyst))["total"] == 714
    started = time.perf_counter()
    report = await validate(client, demo, analyst)
    assert time.perf_counter() - started < 3 and report["seconds"] < 3
    data = (await client.get(f"{base(demo)}/export", params={"kind": "anexo"}, headers=analyst)).content
    started = time.perf_counter()
    job = await analysed(client, demo, analyst, data)
    analysis = time.perf_counter() - started
    assert job["counts"]["rows_unchanged"] == 714 and analysis < 10, analysis
    started = time.perf_counter()
    assert (await apply(client, demo, analyst, job)).status_code == 200
    assert time.perf_counter() - started < 10


async def test_the_analysis_writes_and_locks_nothing(client, demo):
    """E16 fleco 1: analysing a book (a QTY changed, a row removed in Sustituir, a NOC
    unknown to the master) while another session holds the CUF's rows and articles
    locked: it neither waits nor writes, and it still says what the result would be."""
    from integration.test_anexo_xi import _book, _bytes, _changed_demo

    admin = demo["people"]["admin"]
    data, _removed = _changed_demo()
    book = _book(data)
    ws = book["CON_FUN_FIS"]
    radio = next(r for r in ws.iter_rows(min_row=3) if r[4].value == "E12300")
    radio[8].value = "999999X999999"
    data = _bytes(book)
    mid = demo["materials"][TRAMONTANA].id
    async with async_session_factory() as db:
        audits = len((await db.execute(select(AuditLog.id))).all())
        rows_before = sorted((str(i), v) for i, v in (await db.execute(select(PhysicalPosition.id, PhysicalPosition.version).where(PhysicalPosition.material_id == mid))).all())
    async with async_session_factory() as holder:
        await holder.execute(select(PhysicalPosition.id).where(PhysicalPosition.material_id == mid).with_for_update())
        from app.models import Article

        await holder.execute(select(Article.id).with_for_update())
        job = await asyncio.wait_for(analysed(client, demo, admin, data, mode="replace"), timeout=30)
        await holder.rollback()
    report = job["report"]
    assert job["counts"]["rows_modified"] == 2 and job["counts"]["rows_removed"] == 1
    assert report["result"]["provisional_articles"] == [{"noc": "999999X999999", "name": radio[11].value}]
    async with async_session_factory() as db:
        assert len((await db.execute(select(AuditLog.id))).all()) == audits
        rows_after = sorted((str(i), v) for i, v in (await db.execute(select(PhysicalPosition.id, PhysicalPosition.version).where(PhysicalPosition.material_id == mid))).all())
    assert rows_after == rows_before
    # What it saw is what applying does.
    done = await apply(client, demo, admin, job, confirm_name=TRAMONTANA)
    assert done.status_code == 200, done.text


async def test_the_rows_filter_by_rule_and_severity(client, demo):
    """The CONFUNFIS table's filter (E16): a rule's code, or "con errores bloqueantes" / "con avisos"."""
    analyst = demo["people"]["analista"]
    blocking = (await all_rows(client, demo, analyst, severity="bloqueante"))["items"]
    assert sorted(f"{r['sns']} {r['co']}" for r in blocking) == ["A22140 02", "A24110 01", "B31010 01", "B32020 01", "B35010 01", "D53000 01"]
    warned = (await all_rows(client, demo, analyst, severity="aviso"))["items"]
    assert [r["sns"] for r in warned] == ["B33011"]
    assert [r["sns"] for r in (await all_rows(client, demo, analyst, observation="CF-007"))["items"]] == ["A24110"]
    res = await client.get(f"{base(demo)}/rows", params={"severity": "otra"}, headers=analyst)
    assert res.status_code == 422
