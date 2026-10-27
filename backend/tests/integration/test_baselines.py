"""Baselines and change authorizations (E17) over the demo, the whole flow
of the encargo §2.3: corrected and validated, v1 in review (read-only), four
eyes, reviewed, a change refused without an authorization and made with
one (the diff names it), the book import, the delivery of v1 (not the live
data), v2 closes the authorization, a rejected v3. Plus: the snapshot is
immutable, illegal states, the copy of a CUF carries no baseline, a
baseline's validation is never purged."""

import io
from collections import Counter
import os
import sys
import uuid
import zipfile

import openpyxl
import pytest_asyncio
from sqlalchemy import delete, select

from app.db.base import async_session_factory
from app.models import AppParameter, AuditLog, Baseline, User, UserMaterialRole, ValidationRun
from integration.test_anexo_xi import _jobs, analysed, apply  # noqa: F401
from integration.test_catalog import catalog, imported  # noqa: F401
from integration.test_confunfis import SCRIPTS, TRAMONTANA, all_rows, base, demo, find  # noqa: F401
from integration.test_validation import DEMO_FIGURES
from integration.test_validation import url as validation_url
from integration.test_validation import validate

HEADER = "X-Change-Authorization"


def b_url(demo, rest=""):
    return f"/api/materials/{demo['materials'][TRAMONTANA].id}/{rest}"


@pytest_asyncio.fixture
async def corrected(demo):
    """The demo with its 10 blocking findings corrected (scripts/demo_corregir_hallazgos.py)."""
    sys.path.insert(0, str(SCRIPTS))
    import demo_corregir_hallazgos

    async with async_session_factory() as session:
        email = (await session.execute(
            select(User.email).join(UserMaterialRole, UserMaterialRole.user_id == User.id)
            .where(UserMaterialRole.material_id == demo["materials"][TRAMONTANA].id, UserMaterialRole.role_code == "analista")
        )).scalars().first()
    os.environ["DEMO_ANALYST_EMAIL"] = email
    try:
        assert await demo_corregir_hallazgos.run(set()) == 0
        assert await demo_corregir_hallazgos.run(set()) == 0  # idempotent
    finally:
        os.environ.pop("DEMO_ANALYST_EMAIL", None)
    yield demo


async def _qty(client, demo, headers, row, qty, authorization=None):
    extra = {HEADER: authorization} if authorization else {}
    return await client.patch(f"{base(demo)}/rows/{row['id']}", json={"version": row["version"], "qty": qty}, headers={**headers, **extra})


async def _row(client, demo, headers, sns, co="01", al="00"):
    [row] = find((await all_rows(client, demo, headers))["items"], sns, co, al)
    return row


async def _create(client, demo, headers):
    res = await client.post(b_url(demo, "baselines"), headers=headers)
    assert res.status_code == 201, res.text
    return res.json()


async def _review(client, demo, headers, baseline, approve=True, comment=None):
    action = "approve" if approve else "reject"
    return await client.post(b_url(demo, f"baselines/{baseline['id']}/{action}"), json={"version": baseline["version"], "comment": comment}, headers=headers)


async def test_the_whole_flow(client, corrected):
    demo = corrected
    analyst, reviewer, observer = demo["people"]["analista"], demo["people"]["revisor"], demo["people"]["observador"]
    # 1. Corrected and validated: no blocking finding.
    first = await validate(client, demo, analyst)
    assert first["blocking"] == 0, [f["rule_code"] for f in first["findings"] if f["severity"] == "bloqueante"]
    assert (await client.get(b_url(demo, "baselines/state"), headers=analyst)).json()["edit"] == "free"

    # 2. v1 in review: the CONFUNFIS is read-only.
    v1 = await _create(client, demo, analyst)
    assert (v1["number"], v1["state"], v1["validation"]["blocking"]) == (1, "en_revision", 0)
    segments = await _row(client, demo, analyst, "A21031")
    refused = await _qty(client, demo, analyst, segments, 4)
    assert (refused.status_code, refused.json()["detail"]["code"]) == (409, "baseline_in_review")
    assert (await client.post(b_url(demo, "baselines"), headers=analyst)).json()["detail"]["code"] == "baseline_already_in_review"

    # 3. Four eyes: not the analyst, not the observer; the reviewer.
    assert (await _review(client, demo, analyst, v1)).status_code == 403
    assert (await _review(client, demo, observer, v1)).status_code == 403
    approved = await _review(client, demo, reviewer, v1)
    assert approved.status_code == 200 and approved.json()["state"] == "revisada"
    v1 = approved.json()
    state = (await client.get(b_url(demo, "baselines/state"), headers=analyst)).json()
    assert state["edit"] == "authorization" and state["reviewed"]["number"] == 1 and state["changes_since"] == 0

    # 4. A change without an authorization.
    segments = await _row(client, demo, analyst, "A21031")
    refused = await _qty(client, demo, analyst, segments, 4)
    assert (refused.status_code, refused.json()["detail"]["code"]) == (409, "authorization_required")

    # 5. AC-…-001 (LSA): requested, approved by the reviewer, used.
    asked = await client.post(b_url(demo, "authorizations"), json={"kind": "lsa", "description": "Segmentos reforzados", "reason": "4 por pistón"}, headers=analyst)
    assert asked.status_code == 201, asked.text
    ac = asked.json()
    assert ac["code"] == "AC-TRM-4X4-A-001" and ac["state"] == "solicitada"
    assert (await _qty(client, demo, analyst, segments, 4, ac["id"])).json()["detail"]["code"] == "authorization_not_usable"
    own = await client.post(b_url(demo, f"authorizations/{ac['id']}/approve"), json={"version": ac["version"]}, headers=analyst)
    assert own.status_code == 403
    ac = (await client.post(b_url(demo, f"authorizations/{ac['id']}/approve"), json={"version": ac["version"]}, headers=reviewer)).json()
    assert ac["state"] == "aprobada"
    done = await _qty(client, demo, analyst, segments, 4, ac["id"])
    assert done.status_code == 200, done.text
    assert (await _row(client, demo, analyst, "A21031"))["total"] == "24"
    diff = (await client.get(b_url(demo, "baselines/diff"), params={"before": v1["id"], "after": "actual"}, headers=analyst)).json()
    assert [(e["entity"], e["change"], [f["field"] for f in e["fields"]], [a["code"] for a in e["authorizations"]]) for e in diff["entries"]] == [
        ("row", "modified", ["qty"], ["AC-TRM-4X4-A-001"])
    ]
    assert diff["entries"][0]["fields"][0]["before"] == "3" and diff["entries"][0]["fields"][0]["after"] == "4"
    async with async_session_factory() as session:
        logged = (await session.execute(select(AuditLog).where(AuditLog.action == "confunfis.row_updated", AuditLog.entity_id == segments["id"]))).scalars().all()
    assert logged[-1].after["change_authorization"][0]["code"] == "AC-TRM-4X4-A-001"
    detail = (await client.get(f"{base(demo)}/rows/{segments['id']}", headers=analyst)).json()
    assert (detail["history"][0]["authorization"], detail["history"][0]["baseline"]) == ("AC-TRM-4X4-A-001", None)
    changes = (await client.get(b_url(demo, f"authorizations/{ac['id']}/changes"), headers=analyst)).json()
    assert [(c["action"], c["label"].split(" · ")[0]) for c in changes] == [("modificado", "A21031 01 00")]
    xlsx = await client.get(b_url(demo, "baselines/diff"), params={"before": v1["id"], "after": "actual", "format": "xlsx"}, headers=analyst)
    assert xlsx.status_code == 200 and openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.max_row == 2

    # 6. The book: refused without the authorization, applied with it.
    data = (await client.get(f"{base(demo)}/export", params={"kind": "anexo"}, headers=analyst)).content
    job = await analysed(client, demo, analyst, data)
    refused = await apply(client, demo, analyst, job)
    assert (refused.status_code, refused.json()["detail"]["code"]) == (409, "authorization_required")
    job = await analysed(client, demo, analyst, data)
    ok = await client.post(f"{base(demo)}/import/{job['id']}/apply", json={}, headers={**analyst, HEADER: ac["id"]})
    assert ok.status_code == 200, ok.text

    # 7. The delivery is v1 (QTY 3), saying 1 change is not in it.
    delivery = await client.get(f"{base(demo)}/delivery", headers=analyst)
    assert delivery.status_code == 200, delivery.text
    assert delivery.headers["x-changes-not-included"] == "1" and delivery.headers["x-baseline"] == "1"
    z = zipfile.ZipFile(io.BytesIO(delivery.content))
    anexo = openpyxl.load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.startswith("ANEXO")))))
    rows = list(anexo["CON_FUN_FIS"].iter_rows(min_row=3, values_only=True))
    assert next(r for r in rows if r[4] == "A21031")[7] == 3
    assert anexo.sheetnames[0] == "ESTADO DE VALIDACIÓN" and anexo["ESTADO DE VALIDACIÓN"]["B1"].value.startswith("ENTREGA — baseline v1 revisada por")

    # 8. v2: approving it closes AC-…-001; v1 ↔ v2, 1 modification.
    assert (await validate(client, demo, analyst))["blocking"] == 0
    v2 = await _create(client, demo, analyst)
    v2 = (await _review(client, demo, reviewer, v2)).json()
    assert v2["state"] == "revisada"
    [closed] = [a for a in (await client.get(b_url(demo, "authorizations"), headers=analyst)).json() if a["id"] == ac["id"]]
    assert closed["state"] == "cerrada" and closed["closed_in_baseline"] == 2
    diff = (await client.get(b_url(demo, "baselines/diff"), params={"before": v1["id"], "after": v2["id"]}, headers=analyst)).json()
    assert [(e["entity"], e["change"], [a["code"] for a in e["authorizations"]]) for e in diff["entries"]] == [("row", "modified", ["AC-TRM-4X4-A-001"])]
    assert (await client.get(b_url(demo, "baselines/state"), headers=analyst)).json()["changes_since"] == 0
    detail = (await client.get(f"{base(demo)}/rows/{segments['id']}", headers=analyst)).json()
    assert detail["history"][0]["baseline"] == 2
    used = (await client.get(f"/api/articles/{segments['article']['id']}", headers=analyst)).json()["where_used"]
    assert {u["baseline"] for u in used if u.get("material_id") == str(demo["materials"][TRAMONTANA].id)} == {2}

    # 9. A rejected v3 (with a comment): the CONFUNFIS is no longer read-only.
    v3 = await _create(client, demo, analyst)
    assert (await _review(client, demo, reviewer, v3, approve=False)).status_code == 422
    rejected = await _review(client, demo, reviewer, v3, approve=False, comment="Falta revisar los segmentos")
    assert rejected.json()["state"] == "rechazada"
    segments = await _row(client, demo, analyst, "A21031")
    assert (await _qty(client, demo, analyst, segments, 5)).json()["detail"]["code"] == "authorization_required"


async def test_the_snapshot_is_immutable_and_articles_show_in_the_diff(client, corrected):
    demo = corrected
    analyst, reviewer, admin = demo["people"]["analista"], demo["people"]["revisor"], demo["people"]["admin"]
    await validate(client, demo, analyst)
    v1 = await _create(client, demo, analyst)
    async with async_session_factory() as session:
        before = (await session.get(Baseline, uuid.UUID(v1["id"]))).snapshot_hash
    # An article is global: it changes without an authorization, the baseline does not.
    rod = next(a for a in (await client.get("/api/articles", params={"q": "BIELA"}, headers=admin)).json()["items"])
    changed = await client.patch(f"/api/articles/{rod['id']}", json={"version": rod["version"], "nombre": "BIELA DM-6 REFORZADA"}, headers=admin)
    assert changed.status_code == 200, changed.text
    async with async_session_factory() as session:
        frozen = await session.get(Baseline, uuid.UUID(v1["id"]))
        assert frozen.snapshot_hash == before and frozen.snapshot["articles"][rod["id"]]["nombre"] == "BIELA DM-6"
    diff = (await client.get(b_url(demo, "baselines/diff"), params={"before": v1["id"]}, headers=analyst)).json()
    assert [(e["entity"], e["change"], e["fields"][0]["field"]) for e in diff["entries"]] == [("article", "modified", "nombre")]
    # Four eyes even for an admin: never its own baseline.
    await client.post(b_url(demo, f"baselines/{v1['id']}/withdraw"), json={"version": v1["version"]}, headers=analyst)
    await validate(client, demo, analyst)
    mine = await _create(client, demo, admin)
    assert (await _review(client, demo, admin, mine)).json()["detail"]["code"] == "baseline_review_own"
    assert (await _review(client, demo, reviewer, mine)).status_code == 200


async def test_illegal_states(client, corrected):
    demo = corrected
    analyst, reviewer = demo["people"]["analista"], demo["people"]["revisor"]
    # Not validated (or out of date): no baseline.
    refused = await client.post(b_url(demo, "baselines"), headers=analyst)
    assert (refused.status_code, refused.json()["detail"]["code"]) == (409, "baseline_not_valid")
    await validate(client, demo, analyst)
    v1 = await _create(client, demo, analyst)
    withdrawn = (await client.post(b_url(demo, f"baselines/{v1['id']}/withdraw"), json={"version": v1["version"]}, headers=analyst)).json()
    assert withdrawn["state"] == "retirada"
    late = await _review(client, demo, reviewer, withdrawn)
    assert (late.status_code, late.json()["detail"]["code"]) == (409, "baseline_state_invalid")
    # Withdrawn, the CONFUNFIS is editable again (no reviewed baseline: free).
    segments = await _row(client, demo, analyst, "A21031")
    assert (await _qty(client, demo, analyst, segments, 3)).status_code == 200
    assert (await client.post(b_url(demo, "authorizations"), json={"kind": "nada", "description": "x", "reason": "y"}, headers=analyst)).status_code == 422


async def test_a_copy_has_no_baseline_and_a_baselines_run_is_kept(client, corrected):
    demo = corrected
    analyst, reviewer, admin = demo["people"]["analista"], demo["people"]["revisor"], demo["people"]["admin"]
    await validate(client, demo, analyst)
    v1 = await _create(client, demo, analyst)
    await _review(client, demo, reviewer, v1)
    copy = await client.post(f"/api/materials/{demo['materials'][TRAMONTANA].id}/copy", json={"name": "[test] copia con baseline", "noc": None, "copy_members": False}, headers=admin)
    assert copy.status_code == 201, copy.text
    state = (await client.get(f"/api/materials/{copy.json()['id']}/baselines/state", headers=admin)).json()
    assert state["reviewed"] is None and state["edit"] == "free"
    async with async_session_factory() as session:
        session.add(AppParameter(key="validation", value={"keep_runs": 1}))
        await session.commit()
    try:
        for _ in range(3):
            await validate(client, demo, analyst)
        async with async_session_factory() as session:
            mid = demo["materials"][TRAMONTANA].id
            left = set((await session.execute(select(ValidationRun.id).where(ValidationRun.material_id == mid))).scalars())
        assert uuid.UUID(v1["validation"]["id"]) in left and len(left) == 2
    finally:
        async with async_session_factory() as session:
            await session.execute(delete(AppParameter).where(AppParameter.key == "validation"))
            await session.commit()


async def test_permissions(client, corrected):
    demo = corrected
    observer, outsider = demo["people"]["observador"], demo["people"]["outsider"]
    assert (await client.post(b_url(demo, "baselines"), headers=observer)).status_code == 403
    assert (await client.post(b_url(demo, "authorizations"), json={"kind": "lsa", "description": "x", "reason": "y"}, headers=observer)).status_code == 403
    assert (await client.get(b_url(demo, "baselines"), headers=observer)).status_code == 200
    assert (await client.get(b_url(demo, "baselines"), headers=outsider)).status_code == 404
    assert (await client.get(validation_url(demo), headers=outsider)).status_code == 404


async def test_the_reset_leaves_the_demo_as_it_came(client, corrected):
    """seed_demo_confunfis.py --reset after a rehearsal: no baseline, no authorization, the demo's findings."""
    demo = corrected
    analyst, reviewer = demo["people"]["analista"], demo["people"]["revisor"]
    await validate(client, demo, analyst)
    v1 = await _create(client, demo, analyst)
    await _review(client, demo, reviewer, v1)
    await client.post(b_url(demo, "authorizations"), json={"kind": "lsa", "description": "x", "reason": "y"}, headers=analyst)
    assert await demo["seed"].reset() == 0
    assert await demo["seed"].seed() == 0
    assert (await client.get(b_url(demo, "baselines"), headers=analyst)).json() == []
    assert (await client.get(b_url(demo, "authorizations"), headers=analyst)).json() == []
    found = await validate(client, demo, analyst)
    assert dict(Counter(f["rule_code"] for f in found["findings"])) == DEMO_FIGURES
