"""E12 fleco 3: the same SNS with different descriptions per subcategory.
Juanma's case: A22140 (DEPÓSITO DE EXPANSIÓN) is in 06-01 and 06-02 with
PTE. (a) Only the 06-01 row changes to TBD: the diff shows each
subcategory's row before and after, the SNS-level description "varies by
subcategory" with every value (never one picked silently, HR7), and CN-007.
(b) Both rows change: the after of both rows is TBD."""

import io

import openpyxl

from integration.test_catalog import DEMO, admin, analyse, catalog, imported  # noqa: F401


def with_description(subcategories: set[str], value: str) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(DEMO.read_bytes()))
    ws = wb["CONF_FUN"]
    for r in range(2, ws.max_row + 1):
        if ws[f"F{r}"].value == "A22140" and f"{ws[f'A{r}'].value}{ws[f'C{r}'].value}" in subcategories:
            ws[f"H{r}"] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def _diff(client, headers, data):
    _, before = await imported(client, headers)
    _, after = await imported(client, headers, data)
    res = await client.get("/api/catalog/diff", params={"from_version": before["id"], "to_version": after["id"]}, headers=headers)
    assert res.status_code == 200, res.text
    return {i["sns"]: i for i in res.json()["items"]}


async def test_both_rows_changed(client, admin):  # noqa: F811
    items = await _diff(client, admin, with_description({"0601", "0602"}, "TBD"))
    item = items["A22140"]
    rows = {r["subcategory"]: r for r in item["rows"]}
    assert {s: (r["before"]["description"], r["after"]["description"], r["changed"]) for s, r in rows.items()} == {
        "0601": ("PTE", "TBD", ["description"]),
        "0602": ("PTE", "TBD", ["description"]),
    }
    description = next(c for c in item["changes"] if c["field"] == "description")
    assert (description["before"], description["after"], description.get("varies_after")) == ("PTE", "TBD", None)


async def test_one_row_changed(client, admin):  # noqa: F811
    data = with_description({"0601"}, "TBD")
    job = await analyse(client, admin, data)
    assert job["counts"]["cn_007"] == 1
    assert [w["params"] for w in job["report"]["warnings"] if w["code"] == "cn_007"] == [
        {"sns": "A22140", "values": "06-01: TBD · 06-02: PTE"}
    ]
    items = await _diff(client, admin, data)
    item = items["A22140"]
    rows = {r["subcategory"]: (r["before"]["description"], r["after"]["description"], r["changed"]) for r in item["rows"]}
    assert rows == {"0601": ("PTE", "TBD", ["description"]), "0602": ("PTE", "PTE", [])}
    description = next(c for c in item["changes"] if c["field"] == "description")
    assert description["varies_after"] == [{"subcategory": "0601", "value": "TBD"}, {"subcategory": "0602", "value": "PTE"}]
    assert description.get("varies_before") is None


async def test_the_demo_has_no_cn_007(client, admin):  # noqa: F811
    job = await analyse(client, admin, DEMO.read_bytes())
    assert job["counts"]["cn_007"] == 0


async def test_the_sns_card_shows_the_variation(client, admin):  # noqa: F811
    _, version = await imported(client, admin, with_description({"0601"}, "TBD"))
    card = (await client.get(f"/api/catalog/versions/{version['id']}/sns/A22140", headers=admin)).json()
    assert card["description_varies"] == [{"subcategory": "0601", "value": "TBD"}, {"subcategory": "0602", "value": "PTE"}]
    plain = (await client.get(f"/api/catalog/versions/{version['id']}/sns/A21030", headers=admin)).json()
    assert plain["description_varies"] is None
