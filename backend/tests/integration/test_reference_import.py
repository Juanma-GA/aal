"""Excel export and import of a reference table (E03), end to end through
the API: the round trip of D155, two edited descriptions, analysing writes
nothing, errors block applying, a record changed after the analysis stops
it, and the file checks answer with their code."""

import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from sqlalchemy import select, update

from app.db.base import async_session_factory
from app.models import AuditLog, ImportJob, ReferenceTable, ReferenceValue
from excel_compare import workbook_differences
from helpers import bearer, login
from integration.reference_restore import restore_tables  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def admin(client, factory):
    token, _ = await login(client, await factory.user(is_admin=True))
    return bearer(token)


async def export(client, headers, table):
    res = await client.get(f"/api/reference-tables/{table}/export", headers=headers)
    assert res.status_code == 200 and res.headers["content-type"] == XLSX
    assert res.headers["content-disposition"] == f'attachment; filename="{table}.xlsx"'
    return res.content


async def analyse(client, headers, table, data, name="tabla.xlsx"):
    res = await client.post(f"/api/reference-tables/{table}/import", files={"file": (name, data, XLSX)}, headers=headers)
    if res.status_code != 202:
        return res
    # The analysis ran in the background, before the test client returned.
    return await client.get(f"/api/import-jobs/{res.json()['id']}", headers=headers)


def edit(data: bytes, changes: dict[str, object], sheet: str) -> bytes:
    """`changes`: cell -> new value."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for cell, value in changes.items():
        wb[sheet][cell] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def snapshot(table):
    async with async_session_factory() as session:
        rows = (await session.execute(select(ReferenceValue).where(ReferenceValue.table_code == table))).scalars()
        return {v.code: (v.description, v.definition, v.attributes, v.is_active, v.version) for v in rows}


async def test_round_trip_of_d155_has_no_changes(client, admin):
    data = await export(client, admin, "D155")
    sheet = openpyxl.load_workbook(io.BytesIO(data))["D155"]
    assert [c.value for c in sheet[1]] == ["Código", "Descripción", "Definición", "Activo", "Asterisco en la fuente"]
    assert sheet.max_row == 162 and sheet["A2"].data_type == "s" and sheet["A2"].number_format == "@"
    job = (await analyse(client, admin, "D155", data)).json()
    assert job["status"] == "analyzed"
    assert job["counts"] == {"rows": 161, "new": 0, "modified": 0, "unchanged": 161, "errors": 0, "rows_with_errors": 0, "warnings": 0}
    assert "pending" not in job["report"]
    # Exporting again gives the very same workbook, cell by cell.
    assert workbook_differences(data, await export(client, admin, "D155")) == []


async def test_two_edited_descriptions_are_two_modified_and_applying_writes_them(client, admin, restore_tables):
    await restore_tables("D112")
    data = await export(client, admin, "D112")
    sheet = openpyxl.load_workbook(io.BytesIO(data))["D112"]
    first, second = sheet["A2"].value, sheet["A3"].value
    before_first, before_second = sheet["B2"].value, sheet["B3"].value
    before = await snapshot("D112")
    job = (await analyse(client, admin, "D112", edit(data, {"B2": "Primera editada", "B3": "Segunda editada"}, "D112"))).json()
    assert (job["counts"]["modified"], job["counts"]["unchanged"], job["counts"]["errors"]) == (2, 80, 0)
    assert job["report"]["modified"] == [
        {"row": 2, "key": first, "changes": [{"field": "description", "header": "Descripción", "before": before_first, "after": "Primera editada"}]},
        {"row": 3, "key": second, "changes": [{"field": "description", "header": "Descripción", "before": before_second, "after": "Segunda editada"}]},
    ]
    # Analysing wrote nothing.
    assert await snapshot("D112") == before

    applied = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert applied.status_code == 200 and applied.json()["status"] == "applied"
    after = await snapshot("D112")
    assert after[first][0] == "Primera editada" and after[first][4] == before[first][4] + 1
    assert after[second][0] == "Segunda editada"
    assert {c: v for c, v in after.items() if c not in (first, second)} == {c: v for c, v in before.items() if c not in (first, second)}
    async with async_session_factory() as session:
        actions = (await session.execute(select(AuditLog.action).where(AuditLog.entity_id.in_([job["id"]]) | AuditLog.action.like("reference_value.%")).order_by(AuditLog.created_at.desc()).limit(3))).scalars().all()
    assert sorted(actions) == ["import.applied", "reference_value.updated", "reference_value.updated"]
    again = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert (again.status_code, again.json()["detail"]["code"]) == (409, "import_already_applied")


async def test_new_values_and_a_code_typed_as_a_number(client, admin, restore_tables):
    await restore_tables("D114")
    data = await export(client, admin, "D114")
    # Row 15: a new code; the "02" retyped by a person as the number 2.
    edited = edit(data, {"A15": "77", "B15": "NUEVA", "D15": "S", "E15": "No temporal", "A2": 2}, "D114")
    job = (await analyse(client, admin, "D114", edited)).json()
    assert job["counts"]["new"] == 1 and job["counts"]["errors"] == 0
    assert job["report"]["new"] == [{"row": 15, "key": "77", "values": {"code": "77", "description": "NUEVA", "definition": "", "is_active": "S", "attr:tipo": "No temporal"}}]
    assert [(w["code"], w["params"]) for w in job["report"]["warnings"]] == [("number_read_as_code", {"value": "2", "code": "02"})]
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)).status_code == 200
    async with async_session_factory() as session:
        new = (await session.execute(select(ReferenceValue).where(ReferenceValue.table_code == "D114", ReferenceValue.code == "77"))).scalar_one()
        assert (new.attributes, new.sort_order, new.is_active) == ({"tipo": "no_temporal"}, 14, True)


async def test_errors_block_applying(client, admin):
    data = await export(client, admin, "D112")
    before = await snapshot("D112")
    job = (await analyse(client, admin, "D112", edit(data, {"A2": None, "B3": "x" * 1001, "D4": "quizá"}, "D112"))).json()
    assert [(e["row"], e["column"], e["code"]) for e in job["report"]["errors"]] == [(2, "A", "required"), (3, "B", "too_long"), (4, "D", "not_yes_no")]
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert (res.status_code, res.json()["detail"]["code"]) == (409, "import_has_errors")
    assert await snapshot("D112") == before


async def test_a_value_changed_after_the_analysis_stops_applying(client, admin, restore_tables):
    await restore_tables("NAE")
    data = await export(client, admin, "NAE")
    job = (await analyse(client, admin, "NAE", edit(data, {"B2": "Cambiado en el Excel"}, "NAE"))).json()
    async with async_session_factory() as session:
        await session.execute(update(ReferenceValue).where(ReferenceValue.table_code == "NAE", ReferenceValue.code == "N").values(version=ReferenceValue.version + 1))
        await session.commit()
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert (res.status_code, res.json()["detail"]) == (409, {"code": "import_stale", "row": 2})
    assert (await snapshot("NAE"))["N"][0] == "Normalizado"


@pytest.mark.parametrize(
    "name, data, code",
    [("tabla.xlsm", b"PK", "excel_wrong_extension"), ("tabla.xls", b"x", "excel_wrong_extension"), ("tabla.xlsx", b"no es un zip", "excel_not_xlsx")],
)
async def test_bad_files_are_refused_before_saving(client, admin, name, data, code):
    res = await analyse(client, admin, "D112", data, name=name)
    assert (res.status_code, res.json()["detail"]["code"]) == (422, code)


async def test_a_wrong_sheet_or_header_is_a_blocking_error(client, admin):
    data = await export(client, admin, "D107")
    job = (await analyse(client, admin, "D112", data)).json()
    assert [e["code"] for e in job["report"]["errors"]] == ["sheet_missing"]
    job = (await analyse(client, admin, "D107", edit(data, {"B1": "Nombre"}, "D107"))).json()
    assert [(e["code"], e["column"]) for e in job["report"]["errors"]] == [("header_mismatch", "B")]


async def test_only_an_admin_imports_and_a_job_is_its_owners(client, factory, admin):
    analyst_token, _ = await login(client, await factory.user())
    analyst = bearer(analyst_token)
    data = await export(client, analyst, "NAE")  # exporting is for anyone
    res = await client.post("/api/reference-tables/NAE/import", files={"file": ("t.xlsx", data, XLSX)}, headers=analyst)
    assert (res.status_code, res.json()["detail"]["code"]) == (403, "admin_required")
    job = (await analyse(client, admin, "NAE", data)).json()
    seen = await client.get(f"/api/import-jobs/{job['id']}", headers=analyst)
    assert (seen.status_code, seen.json()["detail"]["code"]) == (404, "import_job_not_found")


async def test_a_stale_analysis_is_reaped(client, admin):
    data = await export(client, admin, "NAE")
    job = (await analyse(client, admin, "NAE", data)).json()
    async with async_session_factory() as session:
        await session.execute(
            update(ImportJob).where(ImportJob.id == job["id"]).values(status="analyzing", created_at=datetime.now(timezone.utc) - timedelta(hours=2))
        )
        await session.commit()
    reaped = (await client.get(f"/api/import-jobs/{job['id']}", headers=admin)).json()
    assert (reaped["status"], reaped["error_code"]) == ("failed", "import_interrupted")
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert res.json()["detail"]["code"] == "import_not_analyzed"


async def test_the_staged_file_is_removed_after_the_analysis(client, admin):
    from app.services.import_jobs import staging_path

    job = (await analyse(client, admin, "NAE", await export(client, admin, "NAE"))).json()
    assert not staging_path(job["id"]).exists()
    async with async_session_factory() as session:
        stored = await session.get(ImportJob, job["id"])
        assert len(stored.file_sha256) == 64 and stored.file_name == "tabla.xlsx"
        assert (await session.get(ReferenceTable, "NAE")).version >= 1
