"""The E13 flecos of the articles: references (one NOC, several P/N: the
NT §9.1.6 silentblock, NOC3 with PN3 and PN6), the Maestro AAL with its
"Referencias" sheet and its exact round trip, the admin's correction of a
field definition (audited, kept by a seed reload, never leaving stored data
invalid) and the revalidation of every article in the background, with what
appeared and disappeared by rule."""

import io

import openpyxl
import pytest_asyncio
from sqlalchemy import delete, select, update

from app.db.base import async_session_factory
from app.models import ArticleFieldDefinition, ArticleRevalidation, AuditLog, ReferenceValue
from app.services import revalidation
from app.services.reference_seed import FIELDS_TABLE, load_reference_tables
from helpers import bearer, login
from test_articles import XLSX, _wipe, imported

NOC3 = "534033X900003"


@pytest_asyncio.fixture
async def people(client, factory):
    await _wipe()
    async with async_session_factory() as session:
        saved = [
            {c.name: getattr(f, c.key) for c in ArticleFieldDefinition.__table__.columns}
            for f in (await session.execute(select(ArticleFieldDefinition))).scalars()
        ]
        units = [
            {c.name: getattr(v, c.key) for c in ReferenceValue.__table__.columns}
            for v in (await session.execute(select(ReferenceValue).where(ReferenceValue.table_code == "D112"))).scalars()
        ]
    material = await factory.material()
    users = {"admin": await factory.user(is_admin=True), "analista": await factory.user()}
    await factory.assign(users["analista"], material, "analista")
    yield {name: bearer((await login(client, u))[0]) for name, u in users.items()}
    await revalidation.drain()
    await _wipe()
    async with async_session_factory() as session:
        for f in saved:
            await session.execute(update(ArticleFieldDefinition).where(ArticleFieldDefinition.prisma_code == f["prisma_code"]).values(**f))
        for v in units:
            await session.execute(update(ReferenceValue).where(ReferenceValue.id == v["id"]).values(**v))
        await session.execute(delete(ArticleRevalidation))
        await session.commit()


async def _article(client, headers, **body):
    res = await client.post("/api/articles", json={"nombre": "SILENTBLOCK", **body}, headers=headers)
    assert res.status_code == 201, res.text
    return res.json()


async def test_one_noc_several_pn(client, people):
    analyst = people["analista"]
    a = await _article(client, analyst, noc=NOC3, pn="PN3", ncage="A1234")
    [principal] = a["references"]
    assert (a["pn"], principal["principal"], principal["origin"]) == ("PN3", True, "manual")

    res = await client.post(f"/api/articles/{a['id']}/references", json={"pn": "PN6", "ncage": "a1234"}, headers=analyst)
    assert res.status_code == 201, res.text
    a = res.json()
    assert [(r["pn"], r["ncage"], r["principal"]) for r in a["references"]] == [("PN3", "A1234", True), ("PN6", "A1234", False)]
    # Found by either P/N; the pair is unique among the live references.
    assert [i["id"] for i in (await client.get("/api/articles", params={"q": "PN6"}, headers=analyst)).json()["items"]] == [a["id"]]
    other = await client.post("/api/articles", json={"nombre": "OTRO", "pn": "PN6", "ncage": "A1234"}, headers=analyst)
    assert (other.status_code, other.json()["detail"]["code"]) == (409, "article_pn_ncage_taken")
    same = await client.post(f"/api/articles/{a['id']}/references", json={"pn": "PN3", "ncage": "A1234"}, headers=analyst)
    assert same.status_code == 409

    # PN6 becomes the principal one: the list's P/N follows.
    pn6 = next(r for r in a["references"] if r["pn"] == "PN6")
    res = await client.patch(f"/api/articles/{a['id']}/references/{pn6['id']}", json={"version": pn6["version"], "principal": True}, headers=analyst)
    assert res.status_code == 200, res.text
    a = res.json()
    assert a["pn"] == "PN6" and [r["pn"] for r in a["references"] if r["principal"]] == ["PN6"]
    stale = await client.patch(f"/api/articles/{a['id']}/references/{pn6['id']}", json={"version": pn6["version"], "pn": "X"}, headers=analyst)
    assert stale.status_code == 409

    # Removed logically; its pair is free again.
    pn3 = next(r for r in a["references"] if r["pn"] == "PN3")
    res = await client.delete(f"/api/articles/{a['id']}/references/{pn3['id']}", params={"version": pn3["version"], "confirm": "true"}, headers=analyst)
    assert res.status_code == 200 and [r["pn"] for r in res.json()["references"]] == ["PN6"]
    await _article(client, analyst, nombre="SILENTBLOCK ANTIGUO", pn="PN3", ncage="A1234")
    async with async_session_factory() as session:
        actions = (await session.execute(select(AuditLog.action).where(AuditLog.entity_id == a["id"]).order_by(AuditLog.created_at))).scalars().all()
    assert actions == ["article.created", "article.reference_added", "article.reference_updated", "article.reference_removed"]


async def test_a_ncage_needs_its_pn_and_observers_cannot_edit(client, people, factory):
    a = await _article(client, people["analista"], pn="PN1")
    res = await client.post(f"/api/articles/{a['id']}/references", json={"pn": " ", "ncage": "A1234"}, headers=people["analista"])
    assert res.status_code == 422
    res = await client.patch(f"/api/articles/{a['id']}", json={"version": a["version"], "pn": None, "ncage": "A1234"}, headers=people["analista"])
    assert res.status_code == 422 and res.json()["detail"]["problems"][0]["code"] == "pn_required", res.text
    observer = await factory.user()
    await factory.assign(observer, await factory.material(), "observador")
    headers = bearer((await login(client, observer))[0])
    assert (await client.post(f"/api/articles/{a['id']}/references", json={"pn": "PN2"}, headers=headers)).status_code == 403


async def test_the_maestro_carries_every_reference_and_round_trips(client, people):
    admin = people["admin"]
    await imported(client, admin)
    rod = (await client.get("/api/articles", params={"q": "BIELA"}, headers=admin)).json()["items"][0]
    rod = (await client.patch(f"/api/articles/{rod['id']}", json={"version": rod["version"], "pn": "B-1", "ncage": "DEMO2"}, headers=admin)).json()
    await client.post(f"/api/articles/{rod['id']}/references", json={"pn": "B-2", "ncage": "DEMO2"}, headers=admin)

    data = (await client.get("/api/articles/export", params={"template": "maestro"}, headers=admin)).content
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Artículo", "Referencias"]
    refs = [tuple(c.value for c in row) for row in book["Referencias"].iter_rows(min_row=2)]
    assert [r[2:5] for r in refs if r[1] == rod["nombre"]] == [("B-1", "DEMO2", "S"), ("B-2", "DEMO2", "N")]

    res = await client.post("/api/articles/import", params={"template": "maestro"}, files={"file": ("MAESTRO.xlsx", data, XLSX)}, headers=admin)
    job = (await client.get(f"/api/import-jobs/{res.json()['id']}", headers=admin)).json()
    assert (job["counts"]["new"], job["counts"]["modified"], job["counts"]["rejected"], job["counts"]["unchanged"]) == (0, 0, 0, 45)

    # A row with another of its references makes that one the principal.
    ws = book["Artículo"]
    pn_col = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "PN")
    row = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, pn_col).value == "B-1")
    ws.cell(row, pn_col).value = "B-2"
    out = io.BytesIO()
    book.save(out)
    await imported(client, admin, out.getvalue(), "maestro")
    after = (await client.get(f"/api/articles/{rod['id']}", headers=admin)).json()
    assert [(r["pn"], r["principal"]) for r in after["references"]] == [("B-2", True), ("B-1", False)]


async def test_the_admin_corrects_a_field_and_the_articles_are_revalidated(client, people):
    admin = people["admin"]
    await imported(client, admin)
    fields = {f["prisma_code"]: f for f in (await client.get("/api/article-fields", headers=admin)).json()}
    weight = fields["caPesoArt"]
    assert weight["importance"] == "Recomendado"

    assert (await client.patch("/api/article-fields/caPesoArt", json={"version": weight["version"], "importance": "Obligatorio"}, headers=people["analista"])).status_code == 403
    res = await client.patch("/api/article-fields/caPesoArt", json={"version": weight["version"], "importance": "Obligatorio"}, headers=admin)
    assert res.status_code == 200, res.text
    assert (res.json()["importance"], res.json()["edited_by_admin"], res.json()["version"]) == ("Obligatorio", True, weight["version"] + 1)
    await revalidation.drain()
    latest = (await client.get("/api/article-revalidations/latest", headers=people["analista"])).json()
    assert (latest["reason"], latest["detail"], latest["status"]) == ("field_definition", {"field": "caPesoArt"}, "done")
    without_weight = (await client.get("/api/articles", params={"rule": "AR-001", "page_size": 1000}, headers=admin)).json()["total"]
    assert latest["report"]["articles"] == 45 and latest["report"]["by_rule"]["AR-001"]["appeared"] > 0
    assert latest["report"]["by_rule"]["AR-001"]["disappeared"] == 0 and without_weight >= 1

    # Back to Recomendado: they disappear.
    res = await client.patch("/api/article-fields/caPesoArt", json={"version": weight["version"] + 1, "importance": "Recomendado"}, headers=admin)
    await revalidation.drain()
    latest = (await client.get("/api/article-revalidations/latest", headers=admin)).json()
    assert latest["report"]["by_rule"]["AR-001"]["appeared"] == 0 and latest["report"]["by_rule"]["AR-001"]["disappeared"] > 0

    # Never shorter than the stored data; the NOC's length is its format.
    name = fields["NomArt"]
    short = await client.patch("/api/article-fields/NomArt", json={"version": name["version"], "length": 5}, headers=admin)
    assert (short.status_code, short.json()["detail"]["code"]) == (409, "article_field_data_does_not_fit")
    assert short.json()["detail"]["count"] > 0 and short.json()["detail"]["names"]
    noc = await client.patch("/api/article-fields/cdART", json={"version": fields["cdART"]["version"], "length": 20}, headers=admin)
    assert noc.status_code == 422
    decimals = await client.patch("/api/article-fields/vaUsoArt", json={"version": fields["vaUsoArt"]["version"], "decimals": 5}, headers=admin)
    assert (decimals.status_code, decimals.json()["detail"]["code"]) == (422, "article_field_decimals_invalid")

    # A seed reload keeps what the admin corrected and reports it.
    coste = fields["vaUsoArt"]
    assert (await client.patch("/api/article-fields/vaUsoArt", json={"version": coste["version"], "length": 7}, headers=admin)).status_code == 200
    async with async_session_factory() as session:
        report = await (await session.connection()).run_sync(load_reference_tables)
        await session.commit()
        stored = (await session.execute(select(ArticleFieldDefinition.length).where(ArticleFieldDefinition.prisma_code == "vaUsoArt"))).scalar_one()
        audited = (await session.execute(select(AuditLog.action).where(AuditLog.entity_type == "article_field"))).scalars().all()
    assert stored == 7
    assert any((d.table, d.code, d.field, d.stored, d.seed) == (FIELDS_TABLE, "vaUsoArt", "length", 7, 5) for d in report.differences)
    assert audited.count("article_field.updated") >= 3


async def test_the_revalidate_button_and_a_reference_table_change(client, people):
    admin = people["admin"]
    await imported(client, admin)
    assert (await client.post("/api/article-revalidations", headers=people["analista"])).status_code == 403
    res = await client.post("/api/article-revalidations", headers=admin)
    assert res.status_code == 202 and res.json()["reason"] == "manual"
    await revalidation.drain()
    done = (await client.get(f"/api/article-revalidations/{res.json()['id']}", headers=admin)).json()
    assert done["status"] == "done" and done["report"]["changed"] == 0
    assert all(v["appeared"] == v["disappeared"] == 0 for v in done["report"]["by_rule"].values())

    # Deactivating a D112 unit the articles use: AR-001 appears for them.
    items = (await client.get("/api/articles", params={"page_size": 1000}, headers=admin)).json()["items"]
    unit = next(i["fields"]["cdUnidPesoArt"] for i in items if i["fields"]["cdUnidPesoArt"])
    [value] = [v for v in (await client.get("/api/reference-tables/D112/values", headers=admin)).json() if v["code"] == unit]
    res = await client.patch(f"/api/reference-tables/D112/values/{value['id']}", json={"version": value["version"], "is_active": False}, headers=admin)
    assert res.status_code == 200, res.text
    await revalidation.drain()
    latest = (await client.get("/api/article-revalidations/latest", headers=admin)).json()
    assert (latest["reason"], latest["detail"]) == ("reference_table", {"table": "D112"})
    assert latest["report"]["by_rule"]["AR-001"]["appeared"] > 0 and latest["report"]["changed"] > 0
    # Only whether a code is active counts: a new description queues nothing.
    count = await _runs()
    await client.patch(f"/api/reference-tables/D112/values/{value['id']}", json={"version": value["version"] + 1, "description": "otra"}, headers=admin)
    assert await _runs() == count


async def _runs() -> int:
    async with async_session_factory() as session:
        return len((await session.execute(select(ArticleRevalidation.id))).all())
