"""The normalized functional configuration (E10) through the API, with
demo/CONF_FUN_DEMO.xlsx: the import report and its warnings, analysing
writes nothing, applying creates a version and keeps the previous one, the
diff, the tree, the search, the SNS card, the export round trip and the
permissions."""

import io
from pathlib import Path

import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy import delete, func, select, update

from app.db.base import async_session_factory
from app.models import AuditLog, CatalogVersion, MaterialCategory, MaterialSubcategory, ModuleCatalog, ReferenceValue, SnsPosition
from excel_compare import workbook_differences
from helpers import bearer, login

DEMO = Path(__file__).resolve().parents[3] / "demo" / "CONF_FUN_DEMO.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# The NT rules applied to the demo (app/services/catalog.py). CN-002 and
# CN-003 are 2 each, not the 1 the encargo expected: B34510 is a CONJUNTO
# by its code (declared SUBCONJUNTO), and D53000's system D50000 is not in
# the file.
EXPECTED = {
    "rows": 101, "sns": 68, "categories": 4, "subcategories": 5, "uncategorized_rows": 1, "elemento_rows": 1,
    "cn_001": 2, "cn_002": 2, "cn_003": 2, "cn_004": 1, "cn_005": 1, "cn_006": 7,
}


@pytest_asyncio.fixture
async def catalog(factory):
    """Each test starts and ends without any catalog (only tests delete)."""

    async def wipe():
        async with async_session_factory() as session:
            await session.execute(delete(CatalogVersion))
            await session.execute(delete(MaterialSubcategory))
            await session.execute(delete(MaterialCategory))
            await session.commit()

    await wipe()
    yield
    await wipe()


@pytest_asyncio.fixture
async def admin(client, factory, catalog):
    token, _ = await login(client, await factory.user(is_admin=True))
    return bearer(token)


async def analyse(client, headers, data, name="CONF_FUN_DEMO.xlsx"):
    res = await client.post("/api/catalog/import", files={"file": (name, data, XLSX)}, headers=headers)
    assert res.status_code == 202, res.text
    return (await client.get(f"/api/import-jobs/{res.json()['id']}", headers=headers)).json()


async def imported(client, headers, data=None):
    job = await analyse(client, headers, data or DEMO.read_bytes())
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=headers)
    assert res.status_code == 200, res.text
    versions = (await client.get("/api/catalog/versions", headers=headers)).json()
    return job, versions[0]


def edited(data: bytes, cell: str, value: str) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb["CONF_FUN"][cell] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def test_the_report_of_the_demo(client, admin):
    job = await analyse(client, admin, DEMO.read_bytes())
    assert job["status"] == "analyzed", job
    counts = job["counts"]
    assert {k: counts[k] for k in EXPECTED} == EXPECTED
    assert (counts["new"], counts["modified"], counts["removed"], counts["errors"]) == (68, 0, 0, 0)
    warnings = job["report"]["warnings"]
    assert [w["params"]["sns"] for w in warnings if w["code"] == "cn_001"] == ["A21031", "B33011"]
    assert [w["params"]["sns"] for w in warnings if w["code"] == "cn_002"] == ["A22200", "B34510"]
    assert [w["params"]["sns"] for w in warnings if w["code"] == "cn_003"] == ["D53000", "B34510"]
    assert [w["params"]["sns"] for w in warnings if w["code"] == "cn_004"] == ["ZDM100"]
    assert [(w["row"], w["params"]["sns"]) for w in warnings if w["code"] == "cn_005"] == [(102, "E11000")]
    marks = [w["params"]["mark"] for w in warnings if w["code"] == "cn_006"]
    assert (marks.count("PTE"), marks.count("TBD")) == (6, 1)
    # Analysing wrote nothing.
    async with async_session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(CatalogVersion)) == 0
        assert await session.scalar(select(func.count()).select_from(MaterialCategory)) == 0


async def test_apply_creates_a_version_and_keeps_the_previous_one(client, admin):
    _job, first = await imported(client, admin)
    assert (first["number"], first["is_active"], first["label"], first["counts"]["sns"]) == (1, True, "CONF_FUN_DEMO", 68)
    # The same file again: nothing new, nothing modified.
    again = await analyse(client, admin, DEMO.read_bytes())
    assert (again["counts"]["new"], again["counts"]["modified"], again["counts"]["unchanged"], again["counts"]["removed"]) == (0, 0, 68, 0)
    # One name changed in the Excel: 1 modified, with its diff.
    # (A21010 is also in 0602 with its old name, which stays as an extra one.)
    changed = edited(DEMO.read_bytes(), "G6", "CIGÜEÑAL FORJADO")
    job = await analyse(client, admin, changed)
    name_change = {"field": "name", "header": "NOMBRE", "before": "CIGÜEÑAL", "after": "CIGÜEÑAL FORJADO"}
    extra_change = {"field": "extra_names", "header": "NOMBRES ADICIONALES", "before": "", "after": "CIGÜEÑAL"}
    [modified] = job["report"]["modified"]
    assert (modified["row"], modified["key"]) == (6, "A21010")
    assert modified["changes"][:2] == [name_change, extra_change] and [c["field"] for c in modified["changes"]] == ["name", "extra_names", "rows"]
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)).status_code == 200
    versions = (await client.get("/api/catalog/versions", headers=admin)).json()
    assert [(v["number"], v["is_active"]) for v in versions] == [(2, True), (1, False)]
    # The old name only lives on in 0602: an extra name, not a CN-001.
    card = (await client.get(f"/api/catalog/versions/{versions[0]['id']}/sns/A21010", headers=admin)).json()
    assert card["extra_names"] == [{"name": "CIGÜEÑAL", "origin": "OTRA_SUBCATEGORIA"}] and versions[0]["counts"]["cn_001"] == 2
    async with async_session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(SnsPosition)) == 136  # both kept
        actions = (await session.execute(select(AuditLog.action).where(AuditLog.entity_id.in_([v["id"] for v in versions])))).scalars().all()
    assert actions == ["catalog_version.imported", "catalog_version.imported"]
    # The diff between them.
    diff = (await client.get("/api/catalog/diff", params={"from_version": versions[1]["id"], "to_version": versions[0]["id"]}, headers=admin)).json()
    assert diff["summary"] == {"new": 0, "modified": 1, "removed": 0, "unchanged": 67}
    assert [(i["sns"], i["change"], i["name"]) for i in diff["items"]] == [("A21010", "modified", "CIGÜEÑAL FORJADO")]
    assert diff["items"][0]["changes"] == modified["changes"]


async def test_a_version_imported_meanwhile_stops_an_older_analysis(client, admin):
    stale = await analyse(client, admin, DEMO.read_bytes())
    await imported(client, admin)
    res = await client.post(f"/api/import-jobs/{stale['id']}/apply", headers=admin)
    assert (res.status_code, res.json()["detail"]["code"]) == (409, "catalog_version_changed")


async def test_removed_sns_and_blocking_errors(client, admin):
    await imported(client, admin)
    wb = openpyxl.load_workbook(io.BytesIO(DEMO.read_bytes()))
    ws = wb["CONF_FUN"]
    ws.delete_rows(101, 1)  # ZDM100
    ws["F3"] = "a21-00"  # not six [A-Z0-9]
    ws["E4"] = "NIVEL RARO"
    ws["A5"] = "6"  # "06": normalized (a warning), not an error
    out = io.BytesIO()
    wb.save(out)
    job = await analyse(client, admin, out.getvalue())
    assert [(e["row"], e["column"], e["code"]) for e in job["report"]["errors"]] == [(3, "F", "sns_invalid"), (4, "E", "not_in_table")]
    assert {"code": "number_read_as_code", "params": {"value": "6", "code": "06"}, "row": 5, "column": "A", "header": "ID_CAT_MAT"} in job["report"]["warnings"]
    assert job["counts"]["removed"] == 1 and job["report"]["removed"][0]["key"] == "ZDM100"
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)
    assert res.json()["detail"]["code"] == "import_has_errors"


async def test_reading_the_catalog(client, factory, admin):
    analyst_token, _ = await login(client, await factory.user())
    analyst = bearer(analyst_token)
    _job, version = await imported(client, admin)
    base = f"/api/catalog/versions/{version['id']}"
    cats = (await client.get(f"{base}/categories", headers=analyst)).json()
    assert [c["code"] for c in cats["categories"]] == ["00", "01", "02", "06"] and cats["uncategorized_rows"] == 1
    land = next(c for c in cats["categories"] if c["code"] == "06")
    assert [(s["code"], s["name"]) for s in land["subcategories"]] == [("0601", "VEHÍCULOS DE RUEDAS"), ("0602", "VEHÍCULOS DE CADENAS")]
    # The tree, level by level.
    roots = (await client.get(f"{base}/tree", params={"subcategory": "0601"}, headers=analyst)).json()
    assert [r["sns"] for r in roots] == ["000000", "A20000", "A70000", "B00000", "B30000", "B34510", "D10000", "D53000"]
    motor = (await client.get(f"{base}/tree", params={"subcategory": "0601", "parent": "A20000"}, headers=analyst)).json()
    assert [(n["sns"], n["child_count"]) for n in motor] == [("A21000", 3), ("A22000", 2), ("A23000", 1), ("A24000", 1), ("A27000", 1)]
    tracks = (await client.get(f"{base}/tree", params={"subcategory": "0602"}, headers=analyst)).json()
    assert "B35000" in [r["sns"] for r in tracks]  # its system is not in 0602
    # The card.
    card = (await client.get(f"{base}/sns/A21031", headers=analyst)).json()
    assert (card["name"], card["level"], card["parent"]["sns"]) == ("SEGMENTOS", "SUBCONJUNTO", "A21030")
    assert card["extra_names"] == [{"name": "JUEGO DE SEGMENTOS", "origin": "CN-001"}]
    assert {a["subcategory_code"] for a in card["appearances"]} == {"0601", "0602"}
    tire = (await client.get(f"{base}/sns/B33011", headers=analyst)).json()
    assert tire["extra_names"] == [{"name": "VÁLVULA DE INFLADO", "origin": "ELEMENTO"}] and tire["declared_levels"] == ["ELEMENTO", "SUBCONJUNTO"]
    radio = (await client.get(f"{base}/sns/E11000", headers=analyst)).json()
    assert radio["description"] == "EQUIPOS DE RADIO DE COMBATE, INCLUIDAS ANTENAS Y CABLEADO DE RF"
    assert (await client.get(f"{base}/sns/A22140", headers=analyst)).json()["description_pending"] is True
    # The search.
    found = (await client.get(f"{base}/search", params={"q": "a27010"}, headers=analyst)).json()
    assert [i["sns"] for i in found["items"]] == ["A27010"]
    pistons = (await client.get(f"{base}/search", params={"q": "A21030"}, headers=analyst)).json()["items"][0]
    assert pistons["subcategories"] == ["0601", "0602"]
    by_extra = (await client.get(f"{base}/search", params={"q": "válvula"}, headers=analyst)).json()
    assert [i["sns"] for i in by_extra["items"]] == ["B33011"]
    paged = (await client.get(f"{base}/search", params={"q": "A2", "page_size": 5, "page": 2}, headers=analyst)).json()
    assert paged["total"] > 10 and len(paged["items"]) == 5
    warnings = (await client.get(f"{base}/warnings", headers=analyst)).json()
    assert warnings["counts"] == {k: EXPECTED[k] for k in ("cn_001", "cn_002", "cn_003", "cn_004", "cn_005", "cn_006")}


async def test_export_round_trip(client, admin):
    _job, version = await imported(client, admin)
    full = await client.get(f"/api/catalog/versions/{version['id']}/export", headers=admin)
    assert full.status_code == 200 and full.headers["content-disposition"] == 'attachment; filename="CONF_FUN_v1.xlsx"'
    sheet = openpyxl.load_workbook(io.BytesIO(full.content))["CONF_FUN"]
    assert [c.value for c in sheet[1]] == ["ID_CAT_MAT", "CAT_MAT", "ID_SUBCAT_MAT", "SUBCAT_MAT", "NIVEL FUNCIONAL", "ID_CON_FUN", "NOMBRE", "DESCRIPCIÓN"]
    assert sheet.max_row == 102 and sheet["A2"].value == "00" and sheet["A2"].data_type == "s"
    # Importing the export: nothing changes, and the split description is
    # already repaired (no CN-005).
    job = await analyse(client, admin, full.content, name="CONF_FUN_v1.xlsx")
    assert (job["counts"]["new"], job["counts"]["modified"], job["counts"]["removed"], job["counts"]["cn_005"]) == (0, 0, 0, 0)
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=admin)).status_code == 200
    second = (await client.get("/api/catalog/versions", headers=admin)).json()[0]
    again = await client.get(f"/api/catalog/versions/{second['id']}/export", headers=admin)
    assert workbook_differences(full.content, again.content) == []
    wheels = await client.get(f"/api/catalog/versions/{version['id']}/export", params={"subcategory": "0601"}, headers=admin)
    rows = list(openpyxl.load_workbook(io.BytesIO(wheels.content))["CONF_FUN"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 50 and {r[2] for r in rows} == {"01"} and {r[0] for r in rows} == {"06"}


async def test_only_an_admin_imports(client, factory, admin):
    analyst_token, _ = await login(client, await factory.user())
    analyst = bearer(analyst_token)
    res = await client.post("/api/catalog/import", files={"file": ("x.xlsx", DEMO.read_bytes(), XLSX)}, headers=analyst)
    assert (res.status_code, res.json()["detail"]["code"]) == (403, "admin_required")
    await imported(client, admin)
    assert (await client.get("/api/catalog/versions", headers=analyst)).status_code == 200


async def test_a_change_in_one_subcategory_row_is_a_change(client, admin):
    """A21030 keeps its main description; only its 0602 row changes."""
    await imported(client, admin)
    wb = openpyxl.load_workbook(io.BytesIO(DEMO.read_bytes()))
    ws = wb["CONF_FUN"]
    row = next(r for r in range(2, ws.max_row + 1) if ws[f"F{r}"].value == "A21030" and ws[f"C{r}"].value == "02")
    ws[f"H{row}"] = "OTRA FUNCIÓN EN CADENAS"
    ws[f"J{row}"] = "dato suelto"
    out = io.BytesIO()
    wb.save(out)
    job = await analyse(client, admin, out.getvalue())
    modified = {m["key"]: [c["field"] for c in m["changes"]] for m in job["report"]["modified"]}
    assert modified == {"A21030": ["rows"]}
    assert {"code": "extra_cells_ignored", "params": {"columns": "J"}, "row": row, "column": None, "header": None} in job["report"]["warnings"]


async def test_an_inactive_level_is_refused_for_a_new_version(client, admin):
    async with async_session_factory() as session:
        await session.execute(update(ReferenceValue).where(ReferenceValue.table_code == "NIVEL_FUNCIONAL", ReferenceValue.code == "ELEMENTO").values(is_active=False))
        await session.commit()
    try:
        job = await analyse(client, admin, DEMO.read_bytes())
        assert [(e["row"], e["code"]) for e in job["report"]["errors"]] == [(99, "inactive_in_table")]
    finally:
        async with async_session_factory() as session:
            await session.execute(update(ReferenceValue).where(ReferenceValue.table_code == "NIVEL_FUNCIONAL", ReferenceValue.code == "ELEMENTO").values(is_active=True))
            await session.commit()


async def test_the_module_is_available_with_its_route(db):
    module = await db.get(ModuleCatalog, "catalogo-normalizado")
    assert (module.status, module.route) == ("disponible", "/configuracion-funcional")
