"""The article master (E12) with demo/FICHERO_CARGA_DEMO.xlsx: the import
report and its AR findings, a second import with nothing to do, the round
trip of both templates cell by cell, "dónde se usa", editing (the generic
validation, AR-006 cleared, AR-008), uniqueness and the key without NOC,
the grid's batch with a version conflict, and the permissions."""

import io
import sys
from pathlib import Path

import openpyxl
import pytest_asyncio
from sqlalchemy import delete, select

from app.db.base import async_session_factory
from app.models import Article, AuditLog
from helpers import bearer, login

DEMO = Path(__file__).resolve().parents[3] / "demo" / "FICHERO_CARGA_DEMO.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# The CUF of the demo materials (demo/LEEME_DATOS_DEMO.md).
CUF = {"VTL 4X4 TRAMONTANA": "232033X900001", "MOTOR DIESEL 6 CIL DM-6 280 CV": "281533X900101", "RADIO VHF RT-500": "582033X901001"}


async def _wipe():
    async with async_session_factory() as session:
        await session.execute(delete(Article))
        await session.commit()


@pytest_asyncio.fixture
async def people(client, factory):
    """No article before or after; an admin, and an analyst, a reviewer and
    an observer of a material."""
    await _wipe()
    material = await factory.material()
    users = {"admin": await factory.user(is_admin=True)}
    for role in ("analista", "revisor", "observador"):
        users[role] = await factory.user()
        await factory.assign(users[role], material, role)
    yield {name: bearer((await login(client, u))[0]) for name, u in users.items()}
    await _wipe()


async def analyse(client, headers, data, template="fichero_carga", name="FICHERO_CARGA_DEMO.xlsx"):
    res = await client.post("/api/articles/import", params={"template": template}, files={"file": (name, data, XLSX)}, headers=headers)
    assert res.status_code == 202, res.text
    job = (await client.get(f"/api/import-jobs/{res.json()['id']}", headers=headers)).json()
    assert job["status"] == "analyzed", job
    return job


async def imported(client, headers, data=None, template="fichero_carga"):
    job = await analyse(client, headers, data or DEMO.read_bytes(), template)
    res = await client.post(f"/api/import-jobs/{job['id']}/apply", headers=headers)
    assert res.status_code == 200, res.text
    return job


async def find(client, headers, q):
    [item] = (await client.get("/api/articles", params={"q": q}, headers=headers)).json()["items"]
    return item


def names(job, code):
    return sorted({i["name"] for i in job["report"]["ar"][code]})


async def test_the_demo_report_and_a_second_import(client, people):
    job = await analyse(client, people["admin"], DEMO.read_bytes())
    c = job["counts"]
    assert (c["rows"], c["new"], c["rejected"], c["errors"], c["corrections"]) == (45, 45, 0, 0, 0)
    assert {k: c[f"ar_{k}"] for k in ("001", "002", "003", "004", "005", "006", "007")} == {
        "001": 1, "002": 1, "003": 1, "004": 0, "005": 0, "006": 1, "007": 1,
    }
    assert names(job, "AR-001") == ["TAPÓN DE LLENADO DEPÓSITO"]
    assert names(job, "AR-002") == ["DEPÓSITO DE EXPANSIÓN"]
    assert names(job, "AR-003") == ["ACOPLAMIENTO VISCOSO"]
    assert names(job, "AR-006") == ["BIELA DM-6"]
    assert names(job, "AR-007") == ["CAJA DE CAMBIOS AUTOMÁTICA 6V"]
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=people["admin"])).status_code == 200
    listed = (await client.get("/api/articles", headers=people["observador"])).json()
    assert listed["total"] == 45
    assert len([i for i in listed["items"] if i["noc"] is None]) == 1
    repairable = (await client.get("/api/articles", params={"repairable": "S"}, headers=people["admin"])).json()["total"]
    with_life = (await client.get("/api/articles", params={"life_control": "S"}, headers=people["admin"])).json()["total"]
    assert (repairable, with_life) == (19, 6)
    assert [i["nombre"] for i in (await client.get("/api/articles", params={"rule": "AR-006"}, headers=people["admin"])).json()["items"]] == ["BIELA DM-6"]
    again = await analyse(client, people["admin"], DEMO.read_bytes())
    assert (again["counts"]["new"], again["counts"]["modified"], again["counts"]["unchanged"]) == (0, 0, 45)


def cells(data: bytes) -> list[list[tuple]]:
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    return [[(c.value, c.data_type) for c in row] for row in ws.iter_rows()]


async def test_round_trip(client, people):
    await imported(client, people["admin"])
    export = await client.get("/api/articles/export", params={"template": "fichero_carga"}, headers=people["observador"])
    assert export.status_code == 200
    # Same sheet, headers, order, values and types as the file it came from.
    wb = openpyxl.load_workbook(io.BytesIO(export.content))
    assert wb.sheetnames == ["Artículo"]
    original = cells(DEMO.read_bytes())
    exported = cells(export.content)
    assert exported[:2] == original[:2]
    by_noc = lambda rows: sorted(rows[2:], key=lambda r: (str(r[0][0]), str(r[1][0])))  # noqa: E731
    assert by_noc(exported) == by_noc(original)
    # The Maestro AAL: export -> import -> export, the same.
    maestro = (await client.get("/api/articles/export", params={"template": "maestro"}, headers=people["admin"])).content
    header = cells(maestro)[1]
    assert [v for v, _t in header[30:]] == ["PN", "NCAGE", "OrigenNombre", "cdUniVidaA2", "caVidaArt2", "cdUniVidaA3", "caVidaArt3"]
    job = await analyse(client, people["admin"], maestro, template="maestro", name="MAESTRO_AAL.xlsx")
    assert (job["counts"]["new"], job["counts"]["modified"], job["counts"]["unchanged"]) == (0, 0, 45)
    again = (await client.get("/api/articles/export", params={"template": "maestro"}, headers=people["admin"])).content
    assert cells(again) == cells(maestro)


async def test_where_used_and_the_summary(client, factory, people):
    await imported(client, people["admin"])
    for name, noc in CUF.items():
        await factory.material(name=f"{name} (prueba)", noc=noc)
    tramontana = await find(client, people["observador"], "232033X900001")
    card = (await client.get(f"/api/articles/{tramontana['id']}", headers=people["admin"])).json()
    assert card["where_used"] == [{"kind": "material_cuf", "material_id": card["where_used"][0]["material_id"], "material_name": "VTL 4X4 TRAMONTANA (prueba)"}]
    # A material the user is not in stays unknown to them (as its 404 does).
    assert (await client.get(f"/api/articles/{tramontana['id']}", headers=people["observador"])).json()["where_used"] == []
    used = 0
    for item in (await client.get("/api/articles", headers=people["admin"])).json()["items"]:
        used += bool((await client.get(f"/api/articles/{item['id']}", headers=people["admin"])).json()["where_used"])
    assert used == 3
    summary = (await client.get("/api/articles/export-summary", headers=people["admin"])).json()
    assert (summary["articles"], summary["with_findings"], summary["by_rule"]["AR-006"]) == (45, 5, 1)


async def test_editing_clears_ar_006(client, people):
    await imported(client, people["admin"])
    rod = await find(client, people["analista"], "BIELA DM-6")
    assert rod["finding_codes"] == ["AR-006"]
    body = {"version": rod["version"], "fields": {"snCtrlVidaArt": "s"}, "life_controls": [{"unit_code": "03", "quantity": 5000}]}
    for who in ("revisor", "observador"):
        assert (await client.patch(f"/api/articles/{rod['id']}", json=body, headers=people[who])).status_code == 403
    res = await client.patch(f"/api/articles/{rod['id']}", json=body, headers=people["analista"])
    assert res.status_code == 200, res.text
    out = res.json()
    assert (out["fields"]["snCtrlVidaArt"], out["life_controls"][0]["unit_code"], out["finding_codes"], out["findings"]) == ("S", "03", [], [])
    # Stale version, a value that cannot be stored, a code not in D114, an unknown field.
    assert (await client.patch(f"/api/articles/{rod['id']}", json=body, headers=people["analista"])).json()["detail"]["code"] == "version_conflict"
    bad = await client.patch(
        f"/api/articles/{rod['id']}",
        json={"version": out["version"], "fields": {"caPesoArt": "123456.789", "snRepaArt": "X", "cdUnidPesoArt": "ZZ", "nope": 1}, "life_controls": [{"unit_code": "99", "quantity": 1}]},
        headers=people["analista"],
    )
    assert bad.status_code == 422
    problems = {(p["field"], p["code"]) for p in bad.json()["detail"]["problems"]}
    assert problems == {("caPesoArt", "too_many_decimals"), ("snRepaArt", "not_a_choice"), ("cdUnidPesoArt", "not_in_table"), ("nope", "unknown_field"), ("cdUniVidaA", "not_in_table")}
    async with async_session_factory() as session:
        actions = (await session.execute(select(AuditLog.action).where(AuditLog.entity_id == rod["id"]).order_by(AuditLog.created_at))).scalars().all()
    assert actions == ["article.created", "article.updated"]
    history = (await client.get(f"/api/articles/{rod['id']}/history", headers=people["observador"])).json()
    assert [h["action"] for h in history] == ["article.updated", "article.created"]


async def test_creating_uniqueness_and_ar_008(client, people):
    analyst = people["analista"]
    res = await client.post("/api/articles", json={"nombre": "Junta tórica", "noc": "1234 33x9 00001", "fields": {"snRepaArt": "N"}}, headers=analyst)
    assert res.status_code == 201, res.text
    created = res.json()
    assert created["noc"] == "123433X900001"
    assert created["fields"]["cdSegArt"] == "U" and created["corrections"] == [{"code": "AR-008", "field": "cdSegArt", "value": "U"}]
    assert "AR-001" in created["finding_codes"]
    taken = await client.post("/api/articles", json={"nombre": "Otra", "noc": "123433X900001"}, headers=analyst)
    assert (taken.status_code, taken.json()["detail"]["code"]) == (409, "article_noc_taken")
    pn = {"nombre": "Racor", "pn": "RC-1", "ncage": "demo1"}
    assert (await client.post("/api/articles", json=pn, headers=analyst)).status_code == 201
    again = await client.post("/api/articles", json={**pn, "nombre": "Racor 2"}, headers=analyst)
    assert again.json()["detail"]["code"] == "article_pn_ncage_taken"
    wrong = await client.post("/api/articles", json={"nombre": "X", "noc": "12", "ncage": "ab"}, headers=analyst)
    # E13: an NCAGE without its P/N has no reference to go in.
    assert {p["code"] for p in wrong.json()["detail"]["problems"]} == {"noc_invalid", "ncage_invalid", "pn_required"}
    assert (await client.post("/api/articles", json={"nombre": "Y"}, headers=people["observador"])).status_code == 403
    # Deleted logically: out of the list, the NOC free again.
    assert (await client.delete(f"/api/articles/{created['id']}", params={"version": created["version"], "confirm": "true"}, headers=analyst)).status_code == 204
    assert (await client.get(f"/api/articles/{created['id']}", headers=analyst)).status_code == 404
    assert (await client.post("/api/articles", json={"nombre": "Junta tórica", "noc": "123433X900001"}, headers=analyst)).status_code == 201


async def test_the_key_without_noc(client, people):
    """A row without NOC matches the article without NOC nor P/N of the same
    name; a NOC badly formed is rejected with its row, the rest applies."""
    await imported(client, people["admin"])
    wb = openpyxl.load_workbook(io.BytesIO(DEMO.read_bytes()))
    ws = wb.worksheets[0]
    row = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value == "TAPÓN DE LLENADO DEPÓSITO")
    ws.cell(row, 2).value = "tapón  de llenado  depósito"
    ws.cell(row, 9).value = 3
    ws.cell(3, 1).value = "23203X"
    out = io.BytesIO()
    wb.save(out)
    job = await analyse(client, people["admin"], out.getvalue())
    assert (job["counts"]["new"], job["counts"]["modified"], job["counts"]["rejected"]) == (0, 1, 1)
    assert job["report"]["rejected"][0]["row"] == 3 and job["report"]["rejected"][0]["issues"][0]["code"] == "noc_invalid"
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=people["admin"])).status_code == 200
    cap = await find(client, people["admin"], "llenado")
    assert (cap["nombre"], cap["fields"]["imEstdArt"]) == ("tapón  de llenado  depósito", "3.00")
    assert (await client.get("/api/articles", headers=people["admin"])).json()["total"] == 45
    # Only an admin imports.
    res = await client.post("/api/articles/import", params={"template": "fichero_carga"}, files={"file": ("x.xlsx", DEMO.read_bytes(), XLSX)}, headers=people["analista"])
    assert res.status_code == 403


async def test_the_grid_batch(client, people):
    await imported(client, people["admin"])
    a, b = await find(client, people["analista"], "BIELA DM-6"), await find(client, people["analista"], "BOMBA DE AGUA DM-6")
    # Someone else saves b meanwhile: its row is a conflict, a is saved.
    other = await client.patch(f"/api/articles/{b['id']}", json={"version": b["version"], "fields": {"snMtN1Art": "S"}}, headers=people["admin"])
    assert other.status_code == 200
    res = await client.post(
        "/api/articles/batch",
        json={"items": [{"id": a["id"], "version": a["version"], "fields": {"snMtN4Art": "S"}}, {"id": b["id"], "version": b["version"], "fields": {"snMtN4Art": "S"}}]},
        headers=people["analista"],
    )
    assert res.status_code == 200
    [ok, conflict] = res.json()
    assert ok["ok"] and ok["article"]["fields"]["snMtN4Art"] == "S"
    assert (conflict["ok"], conflict["error"]["code"]) == (False, "version_conflict")
    b_now = (await client.get(f"/api/articles/{b['id']}", headers=people["admin"])).json()
    assert b_now["fields"]["snMtN1Art"] == "S" and b_now["version"] == b["version"] + 1
    assert (await client.post("/api/articles/batch", json={"items": [{"id": a["id"], "version": 1, "fields": {}}]}, headers=people["revisor"])).status_code == 403


async def test_rights(client, people):
    for who, expected in (("admin", (True, True)), ("analista", (True, False)), ("revisor", (False, False)), ("observador", (False, False))):
        r = (await client.get("/api/articles/rights", headers=people[who])).json()
        assert (r["can_edit"], r["can_import"]) == expected, who


async def test_seed_demo_articles_is_idempotent(client, people):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
    import seed_demo_articles as seed

    assert await seed.seed() == 0
    assert (await client.get("/api/articles", headers=people["admin"])).json()["total"] == 45
    writes = select(AuditLog.id).where(AuditLog.action.in_(("article.created", "article.updated")))
    async with async_session_factory() as session:
        first = (await session.execute(writes)).all()
    assert await seed.seed() == 0  # nothing to do the second time
    async with async_session_factory() as session:
        again = (await session.execute(writes)).all()
    assert len(again) == len(first)
    rod = await find(client, people["admin"], "BIELA DM-6")
    assert rod["finding_codes"] == ["AR-006"]


async def test_repeated_imports_keep_life_controls_2_and_3(client, people):
    """A Fichero Carga row without a first pair never moves the second into
    its place: the same file again is "unchanged" and pairs 2-3 stay."""
    await imported(client, people["admin"])
    tramontana = await find(client, people["admin"], "232033X900001")
    body = {"version": tramontana["version"], "life_controls": [{"unit_code": None, "quantity": None}, {"unit_code": "01", "quantity": 7}]}
    res = await client.patch(f"/api/articles/{tramontana['id']}", json=body, headers=people["admin"])
    assert res.status_code == 200, res.text
    wb = openpyxl.load_workbook(io.BytesIO(DEMO.read_bytes()))
    ws = wb.worksheets[0]
    row = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == "232033X900001")
    ws.cell(row, 28).value = None
    ws.cell(row, 29).value = None
    out = io.BytesIO()
    wb.save(out)
    for _ in range(2):
        await imported(client, people["admin"], out.getvalue())
    again = await analyse(client, people["admin"], out.getvalue())
    assert again["counts"]["modified"] == 0
    life = (await client.get(f"/api/articles/{tramontana['id']}", headers=people["admin"])).json()["life_controls"]
    assert [(c["position"], c["unit_code"]) for c in life] == [(1, None), (2, "01")]


async def test_pn_ncage_is_unique_on_import_and_a_null_name_is_422(client, people):
    await imported(client, people["admin"])
    rod = await find(client, people["admin"], "BIELA DM-6")
    assert (await client.patch(f"/api/articles/{rod['id']}", json={"version": rod["version"], "pn": "B-1", "ncage": "DEMO2"}, headers=people["admin"])).status_code == 200
    assert (await client.patch(f"/api/articles/{rod['id']}", json={"version": rod["version"] + 1, "nombre": None}, headers=people["admin"])).status_code == 422
    maestro = (await client.get("/api/articles/export", params={"template": "maestro"}, headers=people["admin"])).content
    wb = openpyxl.load_workbook(io.BytesIO(maestro))
    ws = wb.worksheets[0]
    other = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value == "AMORTIGUADOR")
    ws.cell(other, 31).value, ws.cell(other, 32).value = "B-1", "DEMO2"
    out = io.BytesIO()
    wb.save(out)
    job = await analyse(client, people["admin"], out.getvalue(), template="maestro", name="MAESTRO_AAL.xlsx")
    assert [(r["key"], r["issues"][0]["code"]) for r in job["report"]["rejected"]] == [(ws.cell(other, 1).value, "pn_ncage_taken")]
    assert (await client.post(f"/api/import-jobs/{job['id']}/apply", headers=people["admin"])).status_code == 200
