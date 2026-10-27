"""The validation report (E16) and the ESTADO DE VALIDACIÓN sheet.

- `message()`: a finding in words (es.json: `validation.problems`, the AR
  ones `articles.problems` with the fields by their names, HR21).
- `xlsx()`: a summary by rule and a sheet with every finding and column.
- `pdf()`: what is shown to the client -- the CUF, the traffic light, the
  summary by rule and the detail (reportlab: BSD, pure Python, no system
  libraries; on Windows and Linux alike).
- `status_sheet()`: the first sheet of a working export (Anexo XI, Libro
  AAL, Maestro): "BORRADOR — n errores bloqueantes…" or validated.
"""

import io
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.texts import text
from app.io_excel.writer import excel_text
from app.models import Material, ValidationFinding, ValidationRun
from app.validacion.catalog import BY_CODE, CATALOG

_FONT = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="2F3A4C")
# The traffic light, in the document (not a screen: no theme tokens).
_LIGHT = {"red": "B42318", "amber": "B54708", "green": "067647", "none": "667085"}
_ZONE = ZoneInfo("Europe/Madrid")


def when(at: datetime | None) -> str:
    return at.astimezone(_ZONE).strftime("%d/%m/%Y %H:%M") if at else ""


def message(f: ValidationFinding, field_names: dict[str, str]) -> str:
    if f.rule_code.startswith("AR-"):
        name = lambda code: field_names.get(code, code)  # noqa: E731
        params = {**f.params}
        for key in ("flag", "measure"):
            if isinstance(params.get(key), str):
                params[key] = name(params[key])
        if isinstance(params.get("levels"), str):
            params["levels"] = ", ".join(name(x) for x in params["levels"].split(", "))
        return text(f"articles.problems.{f.problem}", **params, field=name(f.field or ""))
    return text(f"validation.problems.{f.rule_code}.{f.problem}", **f.params)


def _state(run: ValidationRun | None) -> str:
    if run is None:
        return "none"
    return "red" if run.blocking else "amber" if run.warnings or run.accepted else "green"


def _by_rule(found: list[ValidationFinding]) -> list[tuple[str, str, int, int]]:
    """(code, severity, findings, accepted), in the catalog's order, only rules with findings."""
    out = []
    for info in CATALOG:
        mine = [f for f in found if f.rule_code == info.code]
        if mine:
            out.append((info.code, mine[0].severity, len(mine), sum(1 for f in mine if f.acceptance_id)))
    return out


def _header(ws, row: int, headers: tuple[str, ...]) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = _FONT, _FILL


def _put(ws, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value, cell.data_type = excel_text(value), "s"
    else:
        cell.value = value


def xlsx(material: Material, run: ValidationRun, found: list[ValidationFinding], by: str, stale: bool, field_names: dict[str, str],
         justifications: dict[Any, str]) -> bytes:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "RESUMEN"
    lines = [
        (text("validation.report.title"), ""),
        (text("validation.report.cuf"), material.name),
        (text("validation.report.noc"), material.noc or ""),
        (text("validation.report.pn"), material.part_number or ""),
        (text("validation.report.validatedAt"), when(run.created_at)),
        (text("validation.report.validatedBy"), by),
        (text("validation.report.state"), text(f"validation.state.{_state(run)}")),
        ("", text("validation.report.counts", blocking=run.blocking, warnings=run.warnings, accepted=run.accepted)),
    ]
    if stale:
        lines.append(("", text("validation.report.outdated")))
    for r, (label, value) in enumerate(lines, start=1):
        _put(ws, r, 1, label)
        _put(ws, r, 2, value)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    start = len(lines) + 2
    _header(ws, start, (text("validation.report.rule"), text("validation.report.ref"), text("validation.report.severity"), text("validation.report.count"), text("validation.report.acceptedCount")))
    for i, (code, severity, n, accepted) in enumerate(_by_rule(found), start=start + 1):
        _put(ws, i, 1, f"{code} · {text(f'validation.rules.{code}.title')}")
        _put(ws, i, 2, BY_CODE[code].ref)
        _put(ws, i, 3, text(f"validation.severity.{severity}"))
        _put(ws, i, 4, n)
        _put(ws, i, 5, accepted)
    for col, width in zip("ABCDE", (52, 26, 14, 12, 12), strict=True):
        ws.column_dimensions[col].width = width
    detail = book.create_sheet("HALLAZGOS")
    headers = (
        text("validation.report.rule"), text("validation.report.severity"), text("validation.report.scopeCol"), text("validation.report.location"),
        text("validation.report.message"), text("validation.report.field"), text("validation.report.found"), text("validation.report.accepted"),
        text("validation.report.justification"), text("validation.report.ref"),
    )
    _header(detail, 1, headers)
    for r, f in enumerate(found, start=2):
        values = (
            f.rule_code, text(f"validation.severity.{f.severity}"), text(f"validation.scope.{f.scope}"), f.location, message(f, field_names),
            field_names.get(f.field or "", f.field or ""), f.found or "", text("validation.report.yes") if f.acceptance_id else text("validation.report.no"),
            justifications.get(f.acceptance_id, "") if f.acceptance_id else "", BY_CODE[f.rule_code].ref,
        )
        for c, v in enumerate(values, start=1):
            _put(detail, r, c, v)
    for i, width in enumerate((9, 12, 16, 48, 70, 18, 18, 10, 40, 18), start=1):
        detail.column_dimensions[get_column_letter(i)].width = width
    detail.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pdf(material: Material, run: ValidationRun, found: list[ValidationFinding], by: str, stale: bool, field_names: dict[str, str],
        justifications: dict[Any, str]) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=text("validation.report.title"), author="AAL")
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8, leading=10)
    cell = lambda value: Paragraph(_escape(value), small)  # noqa: E731
    state = _state(run)
    story: list[Any] = [
        Paragraph(_escape(text("validation.report.title")), styles["Title"]),
        Paragraph(_escape(f"{text('validation.report.cuf')}: {material.name} · {text('validation.report.noc')} {material.noc or '—'} · "
                          f"{text('validation.report.pn')} {material.part_number or '—'}"), styles["BodyText"]),
        Paragraph(_escape(f"{text('validation.report.validatedAt')} {when(run.created_at)} · {text('validation.report.validatedBy')} {by}"), styles["BodyText"]),
        Spacer(1, 4 * mm),
    ]
    light = Table([[text(f"validation.state.{state}"), text("validation.report.counts", blocking=run.blocking, warnings=run.warnings, accepted=run.accepted)]],
                  colWidths=[60 * mm, 150 * mm])
    light.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{_LIGHT[state]}")), ("TEXTCOLOR", (0, 0), (0, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 11), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#98A2B3")), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [light]
    if stale:
        story += [Spacer(1, 2 * mm), Paragraph(_escape(text("validation.report.outdated")), styles["BodyText"])]
    story += [Spacer(1, 5 * mm), Paragraph(_escape(text("validation.report.summary")), styles["Heading2"])]
    summary = [[text("validation.report.rule"), text("validation.report.ref"), text("validation.report.severity"), text("validation.report.count"),
                text("validation.report.acceptedCount")]]
    summary += [[cell(f"{code} · {text(f'validation.rules.{code}.title')}"), cell(BY_CODE[code].ref), cell(text(f"validation.severity.{severity}")), str(n), str(a)]
                for code, severity, n, a in _by_rule(found)]
    story.append(_table(summary, [110 * mm, 45 * mm, 30 * mm, 25 * mm, 25 * mm]) if len(summary) > 1 else Paragraph(_escape(text("validation.report.none")), styles["BodyText"]))
    story += [Spacer(1, 5 * mm), Paragraph(_escape(text("validation.report.details")), styles["Heading2"])]
    detail = [[text("validation.report.rule"), text("validation.report.severity"), text("validation.report.location"), text("validation.report.message"),
               text("validation.report.accepted")]]
    for f in found:
        accepted = (text("validation.report.yes") + (f": {justifications.get(f.acceptance_id, '')}" if f.acceptance_id else "")) if f.acceptance_id else ""
        detail.append([cell(f.rule_code), cell(text(f"validation.severity.{f.severity}")), cell(f.location), cell(message(f, field_names)), cell(accepted)])
    if len(detail) > 1:
        story.append(_table(detail, [18 * mm, 22 * mm, 75 * mm, 105 * mm, 50 * mm]))

    def footer(canvas: Any, document: Any) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(document.pagesize[0] - 14 * mm, 8 * mm, text("validation.report.page", page=document.page))

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _escape(value: str) -> str:
    return (value or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _table(rows: list[list[Any]], widths: list[float]) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F3A4C")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F7")]),
    ]))
    return table


def status_line(run: ValidationRun | None, stale: bool) -> str:
    if run is None:
        return text("validation.status.notValidated")
    if stale:
        return text("validation.status.draftOutdated", at=when(run.created_at))
    if run.blocking:
        return text("validation.status.draft", blocking=run.blocking, warnings=run.warnings, at=when(run.created_at))
    return text("validation.status.ready", warnings=run.warnings, at=when(run.created_at))


async def status_of(db: AsyncSession, material: Material) -> str:
    """A CUF's validation state, in one line, for an exported book."""
    from app.services import validation

    run = await validation.latest(db, material)
    stale = await validation.outdated(db, material, run) if run is not None else {"data": False, "rules": False}
    return status_line(run, stale["data"] or stale["rules"])


def status_sheet(book: Any, lines: list[tuple[str, str]]) -> None:
    """ESTADO DE VALIDACIÓN, first sheet of a working export: one line per CUF (name, state)."""
    ws = book.create_sheet(text("validation.status.sheet"), 0)
    for r, (name, state) in enumerate(lines, start=1):
        _put(ws, r, 1, name)
        _put(ws, r, 2, state)
        ws.cell(row=r, column=2).font = Font(bold=True)
    _put(ws, len(lines) + 2, 1, text("validation.status.note"))
    ws.cell(row=len(lines) + 2, column=1).alignment = Alignment(wrap_text=False)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 80
