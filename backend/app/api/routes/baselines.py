"""Baselines and change authorizations of a CUF's CONFUNFIS (E17).

- `/api/materials/{id}/baselines`: the timeline (confunfis.ver), "Crear
  baseline" (baseline.crear: an up-to-date validation free of blocking
  findings, one in review at a time), withdraw it (its author), approve or
  reject it (baseline.revisar, never its author).
- `/baselines/state`: the configuration state (no baseline, vN in review, vN
  reviewed + n changes since, open authorizations).
- `/baselines/diff?before=<id>&after=<id|actual>`: by entity and field, each
  change with its authorization; `&format=xlsx` as Excel.
- `/authorizations`: request one (autorizacion.solicitar), approve or reject
  it (autorizacion.aprobar, never who asked), close it (who asked, or the
  admin), and the changes made with it.

The rule that a reviewed CONFUNFIS only changes with an approved
authorization is app/services/changes.py (on every write); a step the state
does not allow is a 409, four eyes a 403.
"""

import io
import uuid
from datetime import date
from typing import Any

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import current_user, require_material_permission
from app.api.routes.catalog import XLSX
from app.core.errors import error_detail
from app.core.texts import text
from app.db.base import get_db
from app.io_excel.writer import excel_text
from app.models import Baseline, ChangeAuthorization, Material, User
from app.schemas.baselines import (
    AuthorizationChangeOut,
    AuthorizationIn,
    AuthorizationOut,
    BaselineOut,
    ConfigurationState,
    DiffEntry,
    DiffOut,
    ReviewIn,
    RunBrief,
    VersionIn,
)
from app.services import baselines, validation
from app.services.baselines import BaselineError
from app.services.permissions import MaterialAccess
from app.services.snapshots import client_ip

router = APIRouter(prefix="/api/materials/{material_id}", tags=["baselines"])

READ, CREATE, REVIEW = "confunfis.ver", "baseline.crear", "baseline.revisar"
REQUEST, DECIDE = "autorizacion.solicitar", "autorizacion.aprobar"
CURRENT = "actual"


def _fail(exc: BaselineError) -> HTTPException:
    return HTTPException(status_code=exc.status, detail=error_detail(exc.code, **exc.params))


async def _names(db: AsyncSession, ids: set[uuid.UUID | None]) -> dict[uuid.UUID, str]:
    ids = {i for i in ids if i}
    return dict((await db.execute(select(User.id, User.display_name).where(User.id.in_(ids)))).all()) if ids else {}


async def _baseline_out(db: AsyncSession, b: Baseline, names: dict[uuid.UUID, str] | None = None) -> BaselineOut:
    names = names if names is not None else await _names(db, {b.created_by, b.reviewed_by})
    run = await baselines.run_of(db, b)
    brief = None
    if run is not None:
        light = validation.summary(run, None)
        brief = RunBrief(id=run.id, state=light["state"], blocking=run.blocking, warnings=run.warnings, accepted=run.accepted, at=run.created_at)
    return BaselineOut(
        id=b.id, number=b.number, scope=b.scope, state=b.state, created_at=b.created_at, created_by=b.created_by,
        created_by_name=names.get(b.created_by), reviewed_at=b.reviewed_at, reviewed_by_name=names.get(b.reviewed_by),
        review_comment=b.review_comment, validation=brief, rows=len(b.snapshot.get("rows", {})), snapshot_hash=b.snapshot_hash, version=b.version,
    )


async def _authorization_out(db: AsyncSession, a: ChangeAuthorization, names: dict[uuid.UUID, str] | None = None) -> AuthorizationOut:
    names = names if names is not None else await _names(db, {a.requested_by, a.decided_by})
    changes = len(await baselines.changes_of(db, a))
    closed_in = (await db.scalar(select(Baseline.number).where(Baseline.id == a.closed_in_baseline_id))) if a.closed_in_baseline_id else None
    return AuthorizationOut(
        id=a.id, code=a.code, number=a.number, kind=a.kind, description=a.description, reason=a.reason, state=a.state,
        requested_at=a.requested_at, requested_by=a.requested_by, requested_by_name=names.get(a.requested_by), decided_at=a.decided_at,
        decided_by_name=names.get(a.decided_by), decision_comment=a.decision_comment, closed_at=a.closed_at, closed_in_baseline=closed_in,
        changes=changes, version=a.version,
    )


# ── baselines ────────────────────────────────────────────────────────────


@router.get("/baselines", response_model=list[BaselineOut])
async def list_baselines(access: MaterialAccess = Depends(require_material_permission(READ)), db: AsyncSession = Depends(get_db)) -> list[BaselineOut]:
    found = await baselines.baselines_of(db, access.material)
    names = await _names(db, {x for b in found for x in (b.created_by, b.reviewed_by)})
    return [await _baseline_out(db, b, names) for b in found]


@router.get("/baselines/state", response_model=ConfigurationState)
async def configuration_state(access: MaterialAccess = Depends(require_material_permission(READ)), db: AsyncSession = Depends(get_db)) -> ConfigurationState:
    st = await baselines.state(db, access.material)
    return ConfigurationState(
        reviewed=await _baseline_out(db, st["reviewed"]) if st["reviewed"] else None,
        in_review=await _baseline_out(db, st["in_review"]) if st["in_review"] else None,
        changes_since=st["changes_since"], open_authorizations=len(st["open_authorizations"]),
        usable_authorizations=[await _authorization_out(db, a) for a in st["usable_authorizations"]], edit=st["edit"],
    )


@router.post("/baselines", response_model=BaselineOut, status_code=status.HTTP_201_CREATED)
async def create_baseline(
    request: Request, access: MaterialAccess = Depends(require_material_permission(CREATE)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> BaselineOut:
    try:
        made = await baselines.create(db, access.material, user, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _baseline_out(db, made)


@router.post("/baselines/{baseline_id}/withdraw", response_model=BaselineOut)
async def withdraw_baseline(
    baseline_id: uuid.UUID, body: VersionIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(CREATE)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> BaselineOut:
    try:
        b = await baselines.get(db, access.material, baseline_id, lock=True)
        await baselines.withdraw(db, access.material, b, body.version, user, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _baseline_out(db, b)


async def _review(baseline_id: uuid.UUID, body: ReviewIn, request: Request, access: MaterialAccess, user: User, db: AsyncSession, approve: bool) -> BaselineOut:
    try:
        b = await baselines.get(db, access.material, baseline_id, lock=True)
        await baselines.review(db, access.material, b, body.version, user, approve, body.comment, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _baseline_out(db, b)


@router.post("/baselines/{baseline_id}/approve", response_model=BaselineOut)
async def approve_baseline(
    baseline_id: uuid.UUID, body: ReviewIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(REVIEW)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> BaselineOut:
    return await _review(baseline_id, body, request, access, user, db, True)


@router.post("/baselines/{baseline_id}/reject", response_model=BaselineOut)
async def reject_baseline(
    baseline_id: uuid.UUID, body: ReviewIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(REVIEW)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> BaselineOut:
    return await _review(baseline_id, body, request, access, user, db, False)


# ── the diff ─────────────────────────────────────────────────────────────


async def _diff(db: AsyncSession, material: Material, before: uuid.UUID, after: str) -> tuple[DiffOut, list[dict[str, Any]]]:
    try:
        old = await baselines.get(db, material, before)
        if after == CURRENT:
            new_doc, until, after_label = await baselines.snapshot(db, material, book=False), None, text("baselines.current")
        else:
            new = await baselines.get(db, material, uuid.UUID(after))
            new_doc, until, after_label = new.snapshot, new.created_at, f"v{new.number}"
    except ValueError:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=error_detail("baseline_not_found")) from None
    except BaselineError as exc:
        raise _fail(exc) from None
    entries = await baselines.with_authorizations(db, material, baselines.diff(old.snapshot, new_doc), old.created_at, until)
    return DiffOut(before=f"v{old.number}", after=after_label, entries=[DiffEntry(**e) for e in entries]), entries


@router.get("/baselines/diff", response_model=DiffOut)
async def baseline_diff(
    before: uuid.UUID, after: str = Query(default=CURRENT, max_length=40), format: str = Query(default="json", pattern="^(json|xlsx)$"),
    access: MaterialAccess = Depends(require_material_permission(READ)), db: AsyncSession = Depends(get_db),
) -> Any:
    out, entries = await _diff(db, access.material, before, after)
    if format == "json":
        return out
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = text("baselines.diffSheet")
    headers = [text(f"baselines.diffColumns.{c}") for c in ("entity", "label", "change", "field", "before", "after", "authorization")]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F3A4C")
    r = 2
    for e in entries:
        codes = ", ".join(a["code"] for a in e["authorizations"])
        for f in e["fields"] or [{"field": "", "before": "", "after": ""}]:
            values = [text(f"baselines.entities.{e['entity']}"), e["label"], text(f"baselines.changes.{e['change']}"), f["field"], f["before"], f["after"], codes]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c)
                cell.value, cell.data_type = excel_text(str(v)), "s"
            r += 1
    for c, w in zip("ABCDEFG", (16, 44, 14, 24, 30, 30, 22), strict=True):
        ws.column_dimensions[c].width = w
    data = io.BytesIO()
    book.save(data)
    name = f"COMPARACION_{out.before}_{out.after}_{date.today():%Y%m%d}.xlsx".replace(" ", "_")
    return Response(content=data.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ── change authorizations ────────────────────────────────────────────────


@router.get("/authorizations", response_model=list[AuthorizationOut])
async def list_authorizations(access: MaterialAccess = Depends(require_material_permission(READ)), db: AsyncSession = Depends(get_db)) -> list[AuthorizationOut]:
    found = await baselines.authorizations_of(db, access.material)
    names = await _names(db, {x for a in found for x in (a.requested_by, a.decided_by)})
    return [await _authorization_out(db, a, names) for a in found]


@router.post("/authorizations", response_model=AuthorizationOut, status_code=status.HTTP_201_CREATED)
async def request_authorization(
    body: AuthorizationIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(REQUEST)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> AuthorizationOut:
    try:
        made = await baselines.request_authorization(db, access.material, body.kind, body.description, body.reason, user, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _authorization_out(db, made)


async def _decide(authorization_id: uuid.UUID, body: ReviewIn, request: Request, access: MaterialAccess, user: User, db: AsyncSession, approve: bool) -> AuthorizationOut:
    try:
        a = await baselines.get_authorization(db, access.material, authorization_id)
        await baselines.decide_authorization(db, access.material, a, body.version, user, approve, body.comment, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _authorization_out(db, a)


@router.post("/authorizations/{authorization_id}/approve", response_model=AuthorizationOut)
async def approve_authorization(
    authorization_id: uuid.UUID, body: ReviewIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(DECIDE)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> AuthorizationOut:
    return await _decide(authorization_id, body, request, access, user, db, True)


@router.post("/authorizations/{authorization_id}/reject", response_model=AuthorizationOut)
async def reject_authorization(
    authorization_id: uuid.UUID, body: ReviewIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(DECIDE)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> AuthorizationOut:
    return await _decide(authorization_id, body, request, access, user, db, False)


@router.post("/authorizations/{authorization_id}/close", response_model=AuthorizationOut)
async def close_authorization(
    authorization_id: uuid.UUID, body: VersionIn, request: Request,
    access: MaterialAccess = Depends(require_material_permission(READ)), user: User = Depends(current_user), db: AsyncSession = Depends(get_db),
) -> AuthorizationOut:
    """Who asked for it (or the admin) closes it by hand; the service checks who."""
    try:
        a = await baselines.get_authorization(db, access.material, authorization_id)
        await baselines.close_authorization(db, access.material, a, body.version, user, client_ip(request))
    except BaselineError as exc:
        raise _fail(exc) from None
    await db.commit()
    return await _authorization_out(db, a)


@router.get("/authorizations/{authorization_id}/changes", response_model=list[AuthorizationChangeOut])
async def authorization_changes(
    authorization_id: uuid.UUID, access: MaterialAccess = Depends(require_material_permission(READ)), db: AsyncSession = Depends(get_db),
) -> list[AuthorizationChangeOut]:
    try:
        a = await baselines.get_authorization(db, access.material, authorization_id)
    except BaselineError as exc:
        raise _fail(exc) from None
    found = await baselines.changes_of(db, a)
    names = await _names(db, {c.user_id for c in found})
    current = await baselines.snapshot(db, access.material, book=False)
    labels = await baselines.labels_of(db, found, current)
    return [
        AuthorizationChangeOut(entity_type=c.entity_type, entity_id=c.entity_id, action=c.action, label=labels.get(c.entity_id, c.entity_id), at=c.created_at, by=names.get(c.user_id))
        for c in found
    ]

