"""The exact export of a template: one sheet named as the template says
(and, after it, the sheets of other templates: `extra`, E13),
its header rows (two rows: equal consecutive groups merged), the columns in
order, and the cell format that keeps each value as it is:

- text columns (codes included) are text cells with the "@" format, also
  as the column's format, so Excel neither drops a code's zeros nor turns
  a typed code into a number; a text starting with "=" stays text;
- YES_NO as S/N and DATE as AAAAMMDD, both text;
- INTEGER and DECIMAL as numbers ("0" / "0.00" formats).

Excel's escapes: a literal _xHHHH_ in a text is escaped and a control
character XML cannot carry is written as _xHHHH_, so reader.py reads back
exactly the text written.
"""

import io
import re
from collections.abc import Callable, Iterable, Sequence
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.io_excel.cells import display
from app.io_excel.template import ColumnType, Template

TEXT_FORMAT = "@"
_EXCEL_ESCAPE_RE = re.compile(r"_x[0-9A-Fa-f]{4}_")
_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
# The official Fichero Carga's header (header_look "fichero_carga").
_SIGLE_FONT = Font(name="Calibri", size=8, bold=True)
_SIGLE_CODES_FILL = PatternFill("solid", fgColor="FFFFFFFF")


def excel_text(value: str) -> str:
    text = _EXCEL_ESCAPE_RE.sub(lambda m: "_x005F" + m.group(0), value)
    return _XML_ILLEGAL_RE.sub(lambda m: f"_x{ord(m.group(0)):04X}_", text)


def _number_format(col) -> str:
    if col.type == ColumnType.INTEGER:
        return "0"
    if col.type == ColumnType.DECIMAL:
        return "0." + "0" * col.decimals if col.decimals else "0"
    return TEXT_FORMAT


def _group_runs(columns) -> list[tuple[int, int]]:
    """(first, last) column indexes of each run of equal groups."""
    runs: list[tuple[int, int]] = []
    for i, col in enumerate(columns):
        if runs and columns[runs[-1][0]].headers[0] == col.headers[0]:
            runs[-1] = (runs[-1][0], i)
        else:
            runs.append((i, i))
    return runs


def write_workbook(
    template: Template, rows: Iterable[dict[str, Any]], extra: Sequence[tuple[Template, Iterable[dict[str, Any]]]] = (),
    finish: Callable[[openpyxl.Workbook], None] | None = None,
) -> bytes:
    """`rows` are field -> value (as cells.normalize returns them); `finish`
    adds what is not a template's sheet (E16: ESTADO DE VALIDACIÓN)."""
    workbook = openpyxl.Workbook()
    _write_sheet(workbook.active, template, rows)
    for other, other_rows in extra:
        _write_sheet(workbook.create_sheet(), other, other_rows)
    if finish is not None:
        finish(workbook)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _write_sheet(sheet: Any, template: Template, rows: Iterable[dict[str, Any]]) -> None:
    columns = template.written_columns
    sheet.title = template.sheet

    sigle = template.header_look == "fichero_carga"
    for r in range(template.header_rows):
        for i, col in enumerate(columns):
            cell = sheet.cell(row=r + 1, column=i + 1)
            if sigle:
                cell.font = _SIGLE_FONT
                if r == 1:
                    cell.fill = _SIGLE_CODES_FILL
            else:
                cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.value = excel_text(col.headers[r])
            cell.data_type = "s"
    if template.header_rows == 2:
        for first, last in _group_runs(columns):
            if last > first:
                sheet.merge_cells(start_row=1, start_column=first + 1, end_row=1, end_column=last + 1)

    for i, col in enumerate(columns):
        dimension = sheet.column_dimensions[get_column_letter(i + 1)]
        dimension.width = col.width or max(12, len(col.title) + 2)
        dimension.number_format = _number_format(col)

    for row_number, values in enumerate(rows, start=template.first_data_row):
        for i, col in enumerate(columns):
            value = values.get(col.field)
            if value is None:
                continue
            cell = sheet.cell(row=row_number, column=i + 1)
            cell.number_format = _number_format(col)
            if col.type in (ColumnType.INTEGER, ColumnType.DECIMAL):
                cell.value = value
            else:
                # Always text: openpyxl would take a string starting with
                # "=" as a formula.
                cell.value = excel_text(display(col, value))
                cell.data_type = "s"

    if not sigle:
        sheet.freeze_panes = sheet.cell(row=template.first_data_row, column=1)
