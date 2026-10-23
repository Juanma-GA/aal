"""Reading an uploaded workbook, defensively.

check_upload() runs before the file is even saved: only .xlsx (a .xlsm,
.xls or anything else is refused, and so is an .xlsx with macros inside),
a maximum size, and a maximum uncompressed size read from the zip's own
directory, so a zip bomb is never expanded.

read_rows() opens it with openpyxl in read_only mode twice: once for the
values Excel stored (data_only=True: no formula is ever evaluated, nothing
in the file is run) and once to know which cells hold a formula, because a
formula without a stored value reads as empty and that is an error, never
a blank. defusedxml is installed, so openpyxl parses the XML without entity
expansion.

Headers are checked cell by cell against the template; data rows are read
by position, entirely empty rows skipped, every cell normalized
(cells.py). Problems are Issues: an error code, its params and where (row,
column letter, header), which the frontend turns into a sentence.
"""

import io
import zipfile
from dataclasses import dataclass, field
from itertools import zip_longest
from pathlib import PurePath
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.escape import unescape

from app.io_excel.cells import CellError, normalize
from app.io_excel.template import Template


# The XML of an uploaded workbook must be parsed without entity expansion
# (billion laughs): openpyxl does so only with defusedxml and without lxml.
# Never a silent degradation (HR7): the app refuses to start otherwise.
if not openpyxl.DEFUSEDXML or openpyxl.LXML:
    raise RuntimeError("openpyxl must parse with defusedxml and without lxml (app/io_excel/reader.py)")


class ExcelFileError(Exception):
    """A file that cannot be imported at all: an error code and params."""

    def __init__(self, code: str, **params: Any):
        super().__init__(code)
        self.code = code
        self.params = params


@dataclass(frozen=True)
class Limits:
    max_bytes: int
    max_uncompressed_bytes: int
    max_rows: int


@dataclass(frozen=True)
class Issue:
    code: str
    params: dict[str, Any] = field(default_factory=dict)
    row: int | None = None
    column: str | None = None  # letter
    header: str | None = None

    def as_json(self) -> dict[str, Any]:
        return {"code": self.code, "params": self.params, "row": self.row, "column": self.column, "header": self.header}


@dataclass
class ParsedRow:
    row: int
    values: dict[str, Any]
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)


@dataclass
class ReadResult:
    rows: list[ParsedRow] = field(default_factory=list)
    # Problems of the file as a whole (headers, sheet).
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)


def check_upload(filename: str | None, data: bytes, limits: Limits) -> None:
    name = PurePath(filename or "").name
    if not name.lower().endswith(".xlsx"):
        raise ExcelFileError("excel_wrong_extension")
    if not data:
        raise ExcelFileError("excel_empty")
    if len(data) > limits.max_bytes:
        raise ExcelFileError("excel_too_large", max_mb=limits.max_bytes // (1024 * 1024))
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile:
        raise ExcelFileError("excel_not_xlsx") from None
    names = {e.filename for e in entries}
    if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise ExcelFileError("excel_not_xlsx")
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise ExcelFileError("excel_has_macros")
    if sum(e.file_size for e in entries) > limits.max_uncompressed_bytes:
        raise ExcelFileError("excel_too_large_uncompressed", max_mb=limits.max_uncompressed_bytes // (1024 * 1024))


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    return unescape(str(value)).strip()


def _check_headers(template: Template, header_rows: list[tuple[Any, ...]], result: ReadResult) -> None:
    width = len(template.columns)
    for r in range(template.header_rows):
        cells = header_rows[r] if r < len(header_rows) else ()
        for i, col in enumerate(template.columns):
            found = _header_text(cells[i] if i < len(cells) else None)
            expected = col.headers[r]
            merged = found == "" and (col.extra or (i > 0 and template.columns[i - 1].headers[r] == expected))
            if found.casefold() != expected.casefold() and not merged:
                result.errors.append(
                    Issue("header_mismatch", {"expected": expected, "found": found}, row=r + 1, column=get_column_letter(i + 1))
                )
        for i in range(width, len(cells)):
            extra = _header_text(cells[i])
            if extra:
                result.warnings.append(Issue("extra_column_ignored", {"found": extra}, row=r + 1, column=get_column_letter(i + 1)))


def read_rows(template: Template, data: bytes, limits: Limits) -> ReadResult:
    """The rows of `template`'s sheet. A file openpyxl cannot open, or with
    more rows than allowed, raises ExcelFileError."""
    result = ReadResult()
    try:
        values_book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas_book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception:  # openpyxl raises many types for a broken package
        raise ExcelFileError("excel_corrupt") from None
    try:
        if template.sheet not in values_book.sheetnames:
            result.errors.append(Issue("sheet_missing", {"sheet": template.sheet}))
            return result
        width = len(template.columns)
        try:
            values = values_book[template.sheet].iter_rows(values_only=True)
            cells = formulas_book[template.sheet].iter_rows()
            headers = [next(values, ()) for _ in range(template.header_rows)]
            for _ in range(template.header_rows):
                next(cells, ())
            _check_headers(template, headers, result)
            if result.errors:
                return result
            for row_number, (raw_row, formula_row) in enumerate(zip_longest(values, cells, fillvalue=()), start=template.first_data_row):
                raw = list(raw_row[:width]) + [None] * (width - len(raw_row[:width]))
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
                    continue
                if len(result.rows) >= limits.max_rows:
                    raise ExcelFileError("excel_too_many_rows", max=limits.max_rows)
                parsed = _parse_row(template, row_number, raw, formula_row)
                beyond = [i for i, v in enumerate(raw_row[width:], start=width) if v is not None and not (isinstance(v, str) and not v.strip())]
                if beyond:
                    # Never dropped in silence (HR7): the row says where.
                    parsed.warnings.append(
                        Issue("extra_cells_ignored", {"columns": ", ".join(get_column_letter(i + 1) for i in beyond)}, row=row_number)
                    )
                result.rows.append(parsed)
        except ExcelFileError:
            raise
        except Exception:
            raise ExcelFileError("excel_corrupt") from None
    finally:
        values_book.close()
        formulas_book.close()
    if template.unique_keys:
        _check_duplicate_keys(template, result)
    return result


def _parse_row(template: Template, row_number: int, raw: list[Any], formula_row: tuple[Any, ...]) -> ParsedRow:
    parsed = ParsedRow(row=row_number, values={})
    for i, col in enumerate(template.columns):
        letter = get_column_letter(i + 1)
        where = {"row": row_number, "column": letter, "header": col.title}
        is_formula = i < len(formula_row) and getattr(formula_row[i], "data_type", None) == "f"
        if is_formula and raw[i] is None:
            parsed.errors.append(Issue("formula_without_value", **where))
            continue
        if is_formula:
            parsed.warnings.append(Issue("formula_value_used", **where))
        try:
            normalized = normalize(col, raw[i])
        except CellError as exc:
            parsed.errors.append(Issue(exc.code, exc.params, **where))
            continue
        parsed.values[col.field] = normalized.value
        if normalized.warning:
            parsed.warnings.append(Issue(normalized.warning[0], normalized.warning[1], **where))
    return parsed


def row_key(template: Template, values: dict[str, Any]) -> tuple[Any, ...] | None:
    """The row's key, or None when a key cell has no value."""
    key = tuple(values.get(c.field) for c in template.key_columns)
    return None if any(k is None for k in key) else key


def _check_duplicate_keys(template: Template, result: ReadResult) -> None:
    first: dict[tuple[Any, ...], int] = {}
    for row in result.rows:
        key = row_key(template, row.values)
        if key is None:
            continue
        if key in first:
            row.errors.append(Issue("duplicate_key", {"first_row": first[key]}, row=row.row))
        else:
            first[key] = row.row
