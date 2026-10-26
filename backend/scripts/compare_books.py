"""Compares two Excel books sheet by sheet (E15, verify-confunfis-real):
the values of each sheet as a multiset of rows (the order aside), a text
trimmed, a number as a number (1 and 1.0 alike), an empty cell as None.
Prints one JSON line: per sheet of the first book, its rows in each and
the first rows only one of them has.

    python scripts/compare_books.py BOOK_A BOOK_B [--skip SHEET=N ...] [--min-row R]

--skip SHEET=N leaves the first N columns of that sheet's rows out (the
CUF's columns when the books are of two CUF); --min-row R compares from
row R (3: the data, the headers aside).
"""

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from _common import utf8_output


def _value(v):
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _rows(ws, skip: int, min_row: int) -> Counter:
    out: Counter = Counter()
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        values = tuple(_value(v) for v in row)[skip:]
        if any(v is not None for v in values):
            out[values] += 1
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare two Excel books, sheet by sheet")
    parser.add_argument("first", type=Path)
    parser.add_argument("second", type=Path)
    parser.add_argument("--skip", action="append", default=[], metavar="SHEET=N")
    parser.add_argument("--min-row", type=int, default=1)
    args = parser.parse_args()
    a, b = openpyxl.load_workbook(args.first, read_only=True), openpyxl.load_workbook(args.second, read_only=True)
    skip = {k: int(v) for k, v in (item.rsplit("=", 1) for item in args.skip)}
    report = {}
    for name in a.sheetnames:
        if name not in b.sheetnames:
            report[name] = {"missing": True}
            continue
        ra, rb = _rows(a[name], skip.get(name, 0), args.min_row), _rows(b[name], skip.get(name, 0), args.min_row)
        report[name] = {
            "rows_first": sum(ra.values()), "rows_second": sum(rb.values()), "equal": ra == rb,
            "only_first": [list(r) for r in list((ra - rb).elements())[:5]], "only_second": [list(r) for r in list((rb - ra).elements())[:5]],
        }
    print(json.dumps(report, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    utf8_output()
    sys.exit(main())
